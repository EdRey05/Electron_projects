"""End-to-end validation: the whole app pipeline vs. lifelines 0.27.8.

Runs `km_core.py` (the exact code the Electron app calls) for a set of
scenarios, then recomputes the same quantities with real lifelines and
compares them. This covers the subgroup logic too, not just the estimator.

Run with the reference interpreter (the one with lifelines installed):
    ./kmold/Scripts/python.exe tests/validate_pipeline.py
"""
import json
import os
import subprocess
import sys
import tempfile

import numpy as np
import pandas as pd

from lifelines import KaplanMeierFitter

HERE = os.path.dirname(os.path.abspath(__file__))
APP = os.path.dirname(HERE)
CORE = os.path.join(APP, "python-app", "km_core.py")
sys.path.insert(0, os.path.join(APP, "python-app"))

BASE = r"C:\Users\Administrator\Desktop\Github_EdRey05\Streamlit_projects\Test_units\Project 003"

FAILURES = []


def check(name, cond, detail=""):
    if not cond:
        FAILURES.append(f"{name}: {detail}")
    print(f"  [{'PASS' if cond else 'FAIL'}] {name} {detail}")


def run_core(job):
    """Invoke the app's CLI exactly as Electron does."""
    with tempfile.TemporaryDirectory() as td:
        job_path = os.path.join(td, "job.json")
        out_path = os.path.join(td, "out.json")
        with open(job_path, "w", encoding="utf-8") as f:
            json.dump(job, f)
        proc = subprocess.run(
            [sys.executable, CORE, job_path, "--out", out_path],
            capture_output=True, text=True,
        )
        if not os.path.exists(out_path):
            raise RuntimeError(f"core produced no output. stderr:\n{proc.stderr}")
        with open(out_path, encoding="utf-8") as f:
            return json.load(f)


def lifelines_reference(durations, events, label):
    kmf = KaplanMeierFitter()
    kmf.fit(durations=durations, event_observed=events)
    return {
        "n": len(durations),
        "nEvents": int(np.sum(events)),
        "median": kmf.median_survival_time_,
        "timeline": [float(t) for t in kmf.survival_function_.index],
        "survival": [float(v) for v in kmf.survival_function_.iloc[:, 0]],
        "lower": [float(v) for v in kmf.confidence_interval_.iloc[:, 0]],
        "upper": [float(v) for v in kmf.confidence_interval_.iloc[:, 1]],
        "atRisk": {float(t): int(r["at_risk"] - r["removed"])
                   for t, r in kmf.event_table.iterrows()},
        "label": label,
    }


def compare_curve(ours, ref, tag):
    check(f"[{tag}] n", ours["n"] == ref["n"], f"ours={ours['n']} ref={ref['n']}")
    check(f"[{tag}] events", ours["nEvents"] == ref["nEvents"],
          f"ours={ours['nEvents']} ref={ref['nEvents']}")
    check(f"[{tag}] timeline len", len(ours["timeline"]) == len(ref["timeline"]),
          f"ours={len(ours['timeline'])} ref={len(ref['timeline'])}")
    if len(ours["timeline"]) != len(ref["timeline"]):
        return
    check(f"[{tag}] timeline", np.allclose(ours["timeline"], ref["timeline"], atol=1e-12))
    d = np.max(np.abs(np.array(ours["survival"]) - np.array(ref["survival"])))
    check(f"[{tag}] survival", d < 1e-12, f"max diff={d:.3e}")
    dl = np.max(np.abs(np.array(ours["lower"]) - np.array(ref["lower"])))
    du = np.max(np.abs(np.array(ours["upper"]) - np.array(ref["upper"])))
    check(f"[{tag}] CI", max(dl, du) < 1e-10, f"max diff={max(dl, du):.3e}")
    m1, m2 = ours["median"], ref["median"]
    check(f"[{tag}] median", (m1 == m2) or abs(m1 - m2) < 1e-9, f"ours={m1} ref={m2}")


def main():
    print("=" * 72)
    print("Validating the KM Plotter pipeline (km_core.py) against lifelines 0.27.8")
    print("=" * 72)

    syn = pd.read_csv(os.path.join(BASE, "clinical_synthetic.txt"), sep="\t", comment="#")
    syn_path = os.path.join(BASE, "clinical_synthetic.txt")
    t = syn["OS_MONTHS"].to_numpy(dtype=float)
    e = (syn["OS_STATUS"].astype(str) == "Dead").astype(int).to_numpy()

    # ---- 1. no subgroups -------------------------------------------------
    print("\n=== scenario 1: no subgroups ===")
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [], "alpha": 0.05,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    ours = res["result"]["curves"][0]
    compare_curve(ours, lifelines_reference(t, e, "all"), "nosub")

    # ---- 2. tag subgroups ------------------------------------------------
    print("\n=== scenario 2: tag subgroups (S1) ===")
    groups = [
        {"values": ["Treatment_A"], "name": "A"},
        {"values": ["Treatment_B"], "name": "B"},
        {"values": ["Treatment_C"], "name": "C"},
    ]
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [{"source": "clinical", "column": "S1", "mode": "tags", "groups": groups}],
        "alpha": 0.05,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    check("nCurves == 3", res["result"]["nCurves"] == 3, str(res["result"]["nCurves"]))
    for gname, sname in [("A", "Treatment_A"), ("B", "Treatment_B"), ("C", "Treatment_C")]:
        m = (syn["S1"] == sname).to_numpy()
        curve = next(c for c in res["result"]["curves"] if c["label"] == gname)
        compare_curve(curve, lifelines_reference(t[m], e[m], gname), f"tag-{gname}")

    # ---- 3. numeric range subgroups --------------------------------------
    print("\n=== scenario 3: range subgroups (S2) ===")
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [{"source": "clinical", "column": "S2", "mode": "ranges", "groups": [
            {"lo": 100, "hi": 400, "name": "100-400"},
            {"lo": 400, "hi": 800, "name": "400-800"},
        ]}],
        "alpha": 0.05,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    check("nCurves == 2", res["result"]["nCurves"] == 2, str(res["result"]["nCurves"]))
    s2 = syn["S2"].to_numpy(dtype=float)
    for gname, lo, hi in [("100-400", 100, 400), ("400-800", 400, 800)]:
        m = (s2 >= lo) & (s2 < hi)
        curve = next(c for c in res["result"]["curves"] if c["label"] == gname)
        compare_curve(curve, lifelines_reference(t[m], e[m], gname), f"range-{gname}")

    # ---- 4. two variables (cartesian product) ---------------------------
    print("\n=== scenario 4: two variables (S1 x S4) ===")
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [
            {"source": "clinical", "column": "S1", "mode": "tags", "groups": [
                {"values": ["Treatment_A"], "name": "A"},
                {"values": ["Treatment_B"], "name": "B"},
            ]},
            {"source": "clinical", "column": "S4", "mode": "tags", "groups": [
                {"values": ["1"], "name": "low"},
                {"values": ["2"], "name": "mid"},
            ]},
        ],
        "alpha": 0.05,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    check("nCurves == 4", res["result"]["nCurves"] == 4, str(res["result"]["nCurves"]))
    s1 = syn["S1"].to_numpy()
    s4 = syn["S4"].to_numpy().astype(str)
    for curve in res["result"]["curves"]:
        g1, g2 = curve["label"].split(" | ")
        sa = {"A": "Treatment_A", "B": "Treatment_B"}[g1]
        sb = {"low": "1", "mid": "2"}[g2]
        m = (s1 == sa) & (s4 == sb)
        compare_curve(curve, lifelines_reference(t[m], e[m], curve["label"]), f"2var-{curve['label']}")

    # ---- 5. METABRIC (real-world data, NaNs present) --------------------
    print("\n=== scenario 5: METABRIC (real data with NaNs) ===")
    met_path = os.path.join(BASE, "clinical_METABRIC.txt")
    res = run_core({
        "command": "analyze", "clinicalPath": met_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["0:LIVING"], "eventValues": ["1:DECEASED"],
        "variables": [], "alpha": 0.05,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    met = pd.read_csv(met_path, sep="\t", comment="#")
    met = met.loc[met["OS_STATUS"].notna() & met["OS_MONTHS"].notna()]
    ev = met["OS_STATUS"].astype(str).map({"0:LIVING": 0, "1:DECEASED": 1}).astype(int).to_numpy()
    compare_curve(
        res["result"]["curves"][0],
        lifelines_reference(met["OS_MONTHS"].to_numpy(dtype=float), ev, "metabric"),
        "metabric",
    )

    # ---- 6. alpha = 0.01 (99% CI) ---------------------------------------
    print("\n=== scenario 6: 99% confidence interval ===")
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [], "alpha": 0.01,
    })
    check("ok", res.get("ok"), res.get("error", ""))
    kmf = KaplanMeierFitter(alpha=0.01)
    kmf.fit(durations=t, event_observed=e)
    ours = res["result"]["curves"][0]
    dl = np.max(np.abs(np.array(ours["lower"]) - kmf.confidence_interval_.iloc[:, 0].to_numpy()))
    du = np.max(np.abs(np.array(ours["upper"]) - kmf.confidence_interval_.iloc[:, 1].to_numpy()))
    check("[alpha=0.01] CI", max(dl, du) < 1e-9, f"max diff={max(dl, du):.3e}")

    # ---- 7. Excel export -------------------------------------------------
    print("\n=== scenario 7: Excel export ===")
    res = run_core({
        "command": "export", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [], "alpha": 0.05, "filename": "KM_results.xlsx",
    })
    check("ok", res.get("ok"), res.get("error", ""))
    import base64, io
    raw = base64.b64decode(res["base64"])
    check("xlsx magic bytes", raw[:2] == b"PK", f"got {raw[:4]!r}")
    check("xlsx size", len(raw) > 5000, f"{len(raw)} bytes")
    try:
        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(raw))
        check("workbook sheet", len(wb.sheetnames) == 1, str(wb.sheetnames))
        ws = wb[wb.sheetnames[0]]
        check("sheet has event table", ws.cell(row=6, column=2).value is not None)
    except ImportError:
        print("  (openpyxl not installed here — skipping workbook read-back)")

    # ---- 8. error handling ----------------------------------------------
    print("\n=== scenario 8: error handling ===")
    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": [], "eventValues": [],
        "variables": [], "alpha": 0.05,
    })
    check("empty mapping rejected", res.get("ok") is False, str(res.get("error", ""))[:60])

    res = run_core({
        "command": "analyze", "clinicalPath": syn_path,
        "timeColumn": "NOPE", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [], "alpha": 0.05,
    })
    check("missing column rejected", res.get("ok") is False, str(res.get("error", ""))[:60])

    res = run_core({
        "command": "analyze", "clinicalPath": os.path.join(BASE, "does_not_exist.txt"),
        "timeColumn": "OS_MONTHS", "eventColumn": "OS_STATUS",
        "noEventValues": ["Alive"], "eventValues": ["Dead"],
        "variables": [], "alpha": 0.05,
    })
    check("missing file rejected", res.get("ok") is False, str(res.get("error", ""))[:60])

    print("\n" + "=" * 72)
    if FAILURES:
        print(f"FAILED ({len(FAILURES)}):")
        for f in FAILURES:
            print("  -", f)
        return 1
    print("ALL PIPELINE CHECKS PASSED")
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
