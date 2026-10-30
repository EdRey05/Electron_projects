"""Synthetic BATCH DOCX/sequence ZIPs for real OCR and workbook checks."""
import io
from pathlib import Path
import sys
import tempfile
import zipfile
from docx import Document
from docx.shared import Inches
from PIL import Image,ImageDraw,ImageFont

folder=Path(sys.argv[1]);folder.mkdir(parents=True,exist_ok=True)
font_file=next((p for p in [Path('C:/Windows/Fonts/arial.ttf'),Path('C:/Windows/Fonts/segoeui.ttf')] if p.exists()),None)
font=ImageFont.truetype(str(font_file),32) if font_file else ImageFont.load_default(size=32)
small=ImageFont.truetype(str(font_file),20) if font_file else ImageFont.load_default(size=20)
with tempfile.TemporaryDirectory(prefix='gsh-fixture-') as temporary:
    image=Image.new('RGB',(801,241),'white');draw=ImageDraw.Draw(image)
    for x in [0,200,400,600,800]:draw.line((x,0,x,240),fill='black',width=2)
    for y in [0,80,160,240]:draw.line((0,y,800,y),fill='black',width=2)
    draw.text((20,22),'OD measurement',font=small,fill='black')
    for index,label in enumerate(['OD260','OD280','OD260/OD280','ng/uL']):draw.text((index*200+12,108),label,font=small,fill='black')
    for index,value in enumerate(['0.600','0.300','2.000','30.000']):draw.text((index*200+35,179),value,font=font,fill='black')
    image_path=Path(temporary)/'od.png';image.save(image_path)
    gel=Image.new('RGB',(1200,800),'white');ImageDraw.Draw(gel).text((60,100),'Synthetic gel image',font=font,fill='black')
    gel_path=Path(temporary)/'gel.png';gel.save(gel_path)
    for number in [1,2]:
        doc=Document();doc.add_heading('Synthetic BATCH report',0)
        table=doc.add_table(rows=1,cols=2);table.cell(0,0).text='Plasmid No.';table.cell(0,1).text=f'26TEST00{number}'
        if number==1:
            doc.add_picture(str(gel_path),width=Inches(5))
            doc.add_picture(str(image_path),width=Inches(5))
        stream=io.BytesIO();doc.save(stream)
        bbi=f'BATCH 1234-0{number}'
        filename=folder/(bbi+('.zip' if number==1 else '.ZIP'))
        with zipfile.ZipFile(filename,'w',zipfile.ZIP_DEFLATED) as archive:
            archive.writestr(f'package/BATCH_26TEST00{number}.docx',stream.getvalue())
            archive.writestr(f'package/LA0000{number} {bbi}.seq','>synthetic insert\nACGTRYSWKMBDHVN\n')
            archive.writestr(f'package/LA0000{number}_pUC57.seq','>synthetic final vector\nTTACGTRYSWKMBDHVNCC\n')
print(str(folder))
