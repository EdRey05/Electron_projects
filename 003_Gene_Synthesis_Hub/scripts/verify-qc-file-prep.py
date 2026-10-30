"""Execute real upstream QC processing with synthetic local documents/files."""
import hashlib
import io
import json
from pathlib import Path
import sys
import tempfile
import zipfile
from datetime import date
ROOT=Path(__file__).resolve().parents[1]/'python-app'
sys.path.insert(0,str(ROOT))
from processors.qc_file_prep import run
from docx import Document
from PIL import Image
from openpyxl import Workbook

def make_report(work='XTEST001',bbi='BBI-TEST-01',plasmid='2026A001',vector='pUC57'):
    doc=Document();table=doc.add_table(rows=8,cols=4)
    entries=[('BBI ID',bbi,'',''),('Work No.',work,'',''),('Plasmid No.',plasmid,'',''),('Vector',vector,'Delivery date','2026/10/29'),('Gene Name','GeneOne','',''),('Plasmid Qty.','10 µg','10 µg','')]
    for i,row in enumerate(entries):
        for j,value in enumerate(row):table.cell(i,j).text=value
    doc.add_paragraph('Restriction enzyme digestion analysis')
    qc=doc.add_table(rows=4,cols=2);qc.cell(0,0).text='Enzyme';qc.cell(0,1).text='Expected size';qc.cell(1,0).text='EcoRI';qc.cell(1,1).text='1200 bp'
    od=doc.add_table(rows=2,cols=2);od.cell(0,0).text='OD260/OD280';od.cell(1,0).text='1.95'
    output=io.BytesIO();doc.save(output);return output.getvalue()

def make_zip(folder,stem,work='XTEST001',extra=None):
    folder.mkdir(parents=True,exist_ok=True)
    target=folder/(stem+'.zip')
    with zipfile.ZipFile(target,'w',zipfile.ZIP_DEFLATED) as archive:
        archive.writestr(stem+'/BBI-TEST-01 Gene Synthesis Report.docx',make_report(work))
        for name,data in (extra or {}).items():archive.writestr(stem+'/'+name,data)
    return target

def read_report(filename):
    with zipfile.ZipFile(filename) as archive:
        documents=[n for n in archive.namelist() if n.lower().endswith('.docx')]
        assert len(documents)==1,documents
        return Document(io.BytesIO(archive.read(documents[0])))

def main():
 with tempfile.TemporaryDirectory(prefix='gsh-qc-tests-') as temporary:
    base=Path(temporary);exports=base/'exports';exports.mkdir();logs=[];prompts=[]
    def ask(title,message,choices,fields=None):
        prompts.append(title)
        return {'choice':'proceed','data':{field['key']:field.get('value','') for field in fields or []}}
    def invoke(mode,source=None,**kwargs):return run({'workflow':mode,'inputFolder':str(source or ''),'outputFolder':str(exports),**kwargs},ROOT,logs.append,ask)
    source=base/'new';original=make_zip(source,'XTEST001.2026A001');digest=hashlib.sha256(original.read_bytes()).hexdigest()
    result=invoke('inhouse-new',source,renameFromLog=False);report=read_report(result['files'][0]['path']);assert len(report.tables)==2
    assert hashlib.sha256(original.read_bytes()).hexdigest()==digest
    # Shipment log applies the reference Tube ID suffix rule, without changing IDs.
    book=Workbook();sheet=book.active;sheet.title='Label_Output';sheet.append(['BBI ID','Plasmid ID','Tube ID']);sheet.append(['BBI-TEST-01','2026A001','U12345']);shipment=base/'shipment.xlsx';book.save(shipment)
    result=invoke('inhouse-new',source,renameFromLog=True,paths={'Shipment_and_QC_Log':str(shipment)});assert result['files'][0]['name']=='XTEST001 12345.zip'
    old=base/'old';make_zip(old,'XTEST001.2026A001');result=invoke('inhouse-old',old);assert 'BBI-TEST-01 in pUC57 2026-10-29.2026A001' in result['files'][0]['name']
    mapping=base/'rename.txt';mapping.write_text('BBI-TEST-01\tBBI-TEST-99\n',encoding='utf-8')
    internal=base/'internal';make_zip(internal,'BBI-TEST-01 QC');result=invoke('internal',internal,replacementFile=str(mapping),removeQcTable=True);assert 'BBI-TEST-99' in result['files'][0]['name'];report=read_report(result['files'][0]['path']);assert report.tables[0].cell(0,1).text=='BBI-TEST-99';assert len(report.tables)==2
    wgk=base/'wgk';make_zip(wgk,'WGK001');result=invoke('wgk',wgk,quantity='25 µg');report=read_report(result['files'][0]['path']);assert report.tables[0].cell(5,2).text.strip()=='25 µg'
    qty=base/'qty.txt';qty.write_text('WGK001\t35 µg\n',encoding='utf-8');result=invoke('wgk',wgk,quantityMode='batch',replacementFile=str(qty));assert read_report(result['files'][0]['path']).tables[0].cell(5,2).text.strip()=='35 µg'
    resuspension=base/'resuspension';make_zip(resuspension,'BATCH001');result=invoke('resuspension',resuspension);doc=Document(result['files'][0]['path']);assert doc.tables[0].cell(1,0).text=='BBI-TEST-01';assert doc.tables[0].cell(1,3).text=='1.95'
    # Gel-positive path uses the reference dated directory and exact IDs.
    today=date.today();gels=base/'gels'/str(today.year)/(chr(ord('S')+today.year-2024)+today.strftime('%m%d'));gels.mkdir(parents=True)
    Image.new('RGB',(600,400),'gray').save(gels/'ZTEST001.2026A001.EcoRI.1200bp.png');gel_source=base/'gel';make_zip(gel_source,'ZTEST001.2026A001','ZTEST001')
    result=invoke('inhouse-new',gel_source,renameFromLog=False,paths={'QC_Gel_Images':str(base/'gels')});assert len(read_report(result['files'][0]['path']).inline_shapes)==1
    batch=base/'batch';insert='ACGTCAGT';final='t'*120+insert.lower()+'c'*120
    make_zip(batch,'BATCH001','XTEST001',{'XTEST001.txt':insert,'XTEST001.SEQ':'^^\n'+insert,'XTEST001+pUC57.txt':final,'XTEST001+pUC57.SEQ':'old','Special Growth Protocol.pdf':b'synthetic PDF placeholder'})
    vector=base/'pUC57.fasta';vector.write_text('>pUC57\nACGTCAGT\n',encoding='utf-8')
    result=invoke('batch-post',batch,vectorFasta=str(vector),odValues=['0.6','0.3','2','30']);assert 'BATCH vector and OD' in prompts
    with zipfile.ZipFile(result['files'][0]['path']) as archive:
        names=archive.namelist();assert any(n.endswith('pUC57.fasta') for n in names);assert not any('Special Growth' in n for n in names)
        flank=next(n for n in names if n.endswith('_insert.txt'));assert insert.lower() in archive.read(flank).decode().lower()
        assert any(n.endswith('XTEST001 BBI-TEST-01.fasta') for n in names)
    # Build-from-scratch uses shipped report/PDF templates and local references.
    sorted_dir=base/'sorted';reference=sorted_dir/'XTEST001.2026A001';reference.mkdir(parents=True);(reference/'XTEST001.txt').write_text('ACGTACGT',encoding='utf-8');(reference/'read.ab1').write_bytes(b'synthetic copied trace; not an AB1 parser test')
    align=base/'align';align.mkdir();(align/'XTEST001.pdf').write_bytes(b'synthetic user alignment PDF');(align/'XTEST001.sqd').write_bytes(b'synthetic user alignment SQD')
    book=Workbook();sheet=book.active;sheet.title='Initiated';sheet.append(['Work No.','BBI ID','Vector','Plasmid No.','Tube ID','Gene Name','Length (bp)','Quantity','Prep']);sheet.append(['XTEST001','BBI-TEST-01','pUC57','2026A001','U12345','GeneOne',8,'10 µg','No digestion']);logfile=base/'jobs.xlsx';book.save(logfile)
    result=invoke('scratch',workNos=['XTEST001'],alignmentFolder=str(align),paths={'Job_Log':str(logfile),'Sorted_Sequencing_Files':str(sorted_dir)})
    with zipfile.ZipFile(result['files'][0]['path']) as archive:
        assert any(n.endswith('INSTRUCTIONS-V4.pdf') for n in archive.namelist());seq=next(n for n in archive.namelist() if n.endswith('XTEST001.seq'));assert archive.read(seq)==b'^^\r\nACGTACGT'
        assert any(n.endswith('OriginalSequencingResults/Alignment for XTEST001.sqd') for n in archive.namelist())
    assert hashlib.sha256(original.read_bytes()).hexdigest()==digest
    hostile=base/'hostile';hostile.mkdir()
    with zipfile.ZipFile(hostile/'bad.zip','w') as archive:archive.writestr('../escape.txt','no')
    try:invoke('resuspension',hostile);raise AssertionError('Unsafe ZIP accepted')
    except ValueError as exc:assert 'Unsafe archive' in str(exc)
    print(json.dumps({'status':'passed','workflows':7,'wordRequired':False,'realUpstreamProcessing':True,'nativePromptsExercised':prompts,'originalZipUnchanged':True,'tests':['new/no-gel','shipment renaming','old naming','internal doc/folder rename','WGK single+batch quantities','resuspension values','dated gel lookup+image','BATCH flanks+FASTA+OD','scratch report+templates+references+alignments','archive traversal guard']}))
if __name__=='__main__':main()
