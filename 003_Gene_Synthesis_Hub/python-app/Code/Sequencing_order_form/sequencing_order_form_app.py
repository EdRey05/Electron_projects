"""
Bio Basic Inc. - Canada
Made by: Eduardo Reyes, Ph.D.
Contact: ed5reyes@outlook

Version: 4.5
Date: Oct 15, 2026
Notes: Fills Sample Type (col B) from the sample ID in col C ('Purified
    PCR product' when the ID carries the PFr or PCR marker, 'Plasmid'
    otherwise) and the working concentration (col F, always '40-60
    ng/uL') alongside the other columns. Cell B18 (merged B18:D18) gets
    the order form ID: 'Q' + run date + order form identifier, where
    the run date comes from the Sequencing_ID/date_ID itself (letter
    prefix = year code, U -> 2026 / T -> 2025, digits = MMDD) and the
    identifier is the number after its dash, or 1 without a dash.
"""

import os
import re
import traceback
import threading
import shutil

import openpyxl

import customtkinter as ctk
from tkinter import filedialog, messagebox


PRIMER_LOG_SHEET = "input addon"


def get_primer_info(primerlog, work_nos, logger=print):
    # Open the primer log and select the "input addon" sheet
    wb = openpyxl.load_workbook(primerlog)
    sheet = wb[PRIMER_LOG_SHEET]

    # 2-step column finding: resolve the Work No. column and the primer
    # block by header name on row 1 first (legacy JobID spellings
    # tolerated); historic fixed positions (col A, primer columns from
    # col Q) only as fallback. Whitespace is normalized on both sides.
    # The fired path is logged.
    from openpyxl.utils import get_column_letter

    def norm(h):
        return re.sub(r"\s+", " ", str(h)).strip().lower() if h is not None else ""

    header_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
    idx = {norm(h): i for i, h in enumerate(header_row) if h is not None}

    work_idx = next(
        (idx[c] for c in ("work no.", "jobid", "job id", "job_id") if c in idx),
        None,
    )
    amp_idx = next(
        (i for key, i in idx.items() if "amplification primers" in key), None
    )
    if work_idx is not None and amp_idx is not None:
        primer_start = amp_idx + 1  # individual primers follow the group header
        logger(
            f"'input addon': columns resolved by name "
            f"(Work No.={get_column_letter(work_idx + 1)}, primer block "
            f"from {get_column_letter(primer_start + 1)} after "
            f"'Amplification primers')"
        )
    else:
        work_idx = 0
        primer_start = 16  # historic: Work No.=A, primers from col Q
        logger(
            "'input addon': header row not recognized; falling back to "
            f"fixed columns (Work No.=A, primers from "
            f"{get_column_letter(primer_start + 1)})"
        )

    # Iterate through the rows and make a list of primers for each Work No.
    primers_dict = {str(w).strip(): [] for w in work_nos}

    for row in sheet.iter_rows(min_row=2, values_only=True):
        if work_idx < len(row) and row[work_idx]:
            work_no_stripped = str(row[work_idx]).strip()
            if work_no_stripped in primers_dict:
                for col in range(primer_start, len(row)):
                    if row[col]:
                        primers_dict[work_no_stripped].append(row[col])

    return primers_dict


def get_sequence_sample_ID(seq_validation_file, work_nos):
    # Open the Sequencing Validation file and select the "Sequencing Pending" sheet
    wb = openpyxl.load_workbook(seq_validation_file)
    sequencing_sheet = wb["Sequencing Pending"]

    # Iterate through the rows and make a list of (Work No., Sample ID) tuples in order
    samples = []
    for row in sequencing_sheet.iter_rows(min_row=2, max_row=98, values_only=True):
        if row[0]:
            work_no_stripped = str(row[0]).strip()
            if work_no_stripped in work_nos:
                samples.append(
                    (work_no_stripped, row[3])
                )  # Column A is Work No., Column D is Vector # / Plasmid No.

    return samples


def fill_order_form(
    empty_form,
    ordered_samples,
    primers_dict,
    output_dir,
    seq_validation_file,
    logger=print,
):
    # Open the empty form
    wb = openpyxl.load_workbook(empty_form)
    sheet = wb.active

    # Column H has a common plate date ID (same as folder name where Template 2 is)
    date_ID = os.path.basename(output_dir)
    total_samples = len(ordered_samples)
    logger(f"Starting to fill form with {total_samples} samples...")

    # The first sample space in the form starts in row 20
    row = 20

    warnings = []
    reported_missing_primers = set()

    # Cell B18 (merged B18:D18) carries the order form reference ID:
    # 'Q' + the run date + the order form identifier for that day (the
    # number after the date_ID dash, or 1 when there is no dash). The
    # run date is built from the date_ID itself: the letter prefix
    # encodes the year (R=2023, S=2024, T=2025, U=2026, one letter per
    # year, same base as the inhouse apps' year-letter scheme) and the
    # digits follow as MMDD, e.g. U0304-1 -> 20260304.
    date_ID = str(date_ID).strip()
    date_code, _, form_number = date_ID.partition("-")
    code_digits = re.sub(r"\D", "", date_code)
    year_code = re.sub(r"[^A-Za-z]", "", date_code)
    # Letter -> year, reversed from the inhouse-app year-letter scheme:
    # base_year/base_letter_ord = 2024/'S'.
    year = (
        2024 + (ord(year_code[0].upper()) - ord("S")) if year_code else None
    )
    if code_digits and year:
        identifier = form_number.strip() or "1"
        order_form_ID = f"Q{year}{code_digits}{identifier}"
        sheet["B18"] = order_form_ID
        logger(f"Order form ID (B18): {order_form_ID}")
    else:
        message = (
            f"Could not derive the order form ID (B18) from "
            f"Sequencing_ID '{date_ID}'. Template value left in place."
        )
        logger(f"WARNING: {message}")
        warnings.append(message)

    for work_no, sequence_sample_ID in ordered_samples:
        if (row - 19) % 10 == 0 or row == 20:
            logger(f"  ... processing sample {row - 19}/{total_samples}")

        primers = primers_dict.get(work_no, [])

        # Notify the user when primers are not found for a Work No. (only once per Work No.)
        if not primers:
            if work_no not in reported_missing_primers:
                message = (
                    f"Work No. {work_no}: No primers found in the primer log. Skipping."
                )
                logger(f"WARNING: {message}")
                warnings.append(f"Work No. {work_no}: No primers found.")
                reported_missing_primers.add(work_no)
            continue

        for primer in primers:
            # Sample type (col B): purified PCR products carry the PFr or
            # PCR marker in the sample ID; everything else is a plasmid.
            sample_ID_text = str(sequence_sample_ID)
            sample_type = (
                "Purified PCR product"
                if ("PFr" in sample_ID_text or "PCR" in sample_ID_text)
                else "Plasmid"
            )
            sheet[f"B{row}"] = sample_type
            sheet[f"C{row}"] = sequence_sample_ID
            sheet[f"D{row}"] = primer
            # Working concentration (col F), always the same value.
            sheet[f"F{row}"] = "40-60 ng/uL"
            sheet[f"G{row}"] = work_no
            sheet[f"H{row}"] = date_ID

            formula_E = f'=CONCATENATE(C{row},"._.",G{row},"._.",D{row},"._.",H{row})'
            sheet[f"E{row}"] = formula_E

            validation_filename = os.path.basename(seq_validation_file)
            dir_I = str(output_dir) + f"/[{validation_filename}]Sequencing Pending"
            formula_I = f"=VLOOKUP(C{row},'{dir_I}'!$D:$K,8,0)"
            sheet[f"I{row}"] = formula_I

            formula_J = f'=CONCATENATE(C{row}," ",I{row})'
            sheet[f"J{row}"] = formula_J

            sheet[f"K{row}"] = primer.split("(*", 1)[0].strip()

            formula_L = f'=CONCATENATE(C{row},"._.",G{row},"._.",K{row},"._.",H{row})'
            sheet[f"L{row}"] = formula_L

            row += 1

    # Save the filled form
    output_file = os.path.join(
        output_dir, f"Sequencing_order_form_{os.path.basename(output_dir)}.xlsx"
    )
    logger(f"Saving filled form to: {output_file}")
    wb.save(output_file)
    logger("Save complete.")

    return warnings, output_file


def process_work_nos(
    work_nos, primerlog, seq_validation_file, empty_form, output_dir, logger=print
):
    # Remove blank spaces from work_nos
    work_nos = [str(work_no).strip() for work_no in work_nos if work_no]
    logger(f"Processing {len(work_nos)} Work No.")

    # Get all primers for each work_no
    logger("Fetching primer information from Primer Log...")
    primers_dict = get_primer_info(primerlog, work_nos)
    logger("Primer information loaded.")

    # Get the Sequence Sample IDs (aka Vector # or Plasmid ID) in order
    logger("Fetching sample IDs from Sequencing Validation file...")
    ordered_samples = get_sequence_sample_ID(seq_validation_file, work_nos)
    logger("Sample IDs loaded.")

    if not any(primers_dict.values()):
        raise ValueError("No primers found for any of the Work No.")

    if not ordered_samples:
        raise ValueError(
            "No Sequence Sample IDs (Vector #) found in Sequencing Validation file"
        )

    # Fill out the empty order form and save it
    logger("Generating order form...")
    warnings, output_file = fill_order_form(
        empty_form,
        ordered_samples,
        primers_dict,
        output_dir,
        seq_validation_file,
        logger=logger,
    )
    logger("Order form generation complete.")
    return warnings, output_file


class OrderFormGUI:
    def __init__(
        self, master, primer_log_path=None, empty_form_path=None, seq_layout_path=None
    ):
        self.master = master
        master.title("Sequencing Order Form")
        master.geometry("550x680")

        self.log_textbox = None
        self.seq_layout_path = seq_layout_path

        self.primerlog_valid = False
        self.empty_form_valid = False
        self.seq_validation_valid = False

        # --- Primer Log (Autodetected) ---
        self.primerlog_frame = ctk.CTkFrame(master, fg_color="transparent")
        self.primerlog_frame.pack(pady=(15, 10), padx=15, fill="x")

        ctk.CTkLabel(
            self.primerlog_frame,
            text="Primer Log",
            font=ctk.CTkFont(size=18, weight="bold"),
        ).pack(anchor="center")

        entry_frame_pl = ctk.CTkFrame(self.primerlog_frame, fg_color="transparent")
        entry_frame_pl.pack(fill="x", pady=(2, 2))

        self.entry_primerlog = ctk.CTkEntry(entry_frame_pl, border_width=1)
        self.entry_primerlog.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.button_primerlog = ctk.CTkButton(
            entry_frame_pl,
            text="Select",
            command=self.browse_primerlog,
            width=90,
        )
        self.button_primerlog.pack(side="right")

        self.primerlog_validation_label = ctk.CTkLabel(
            self.primerlog_frame,
            text="",
            font=ctk.CTkFont(size=14),
            anchor="w",
            justify="left",
        )
        self.primerlog_validation_label.pack(fill="x")

        # --- Empty Order Form (Autodetected) ---
        self.empty_form_frame = ctk.CTkFrame(master, fg_color="transparent")
        self.empty_form_frame.pack(pady=(15, 10), padx=15, fill="x")

        ctk.CTkLabel(
            self.empty_form_frame,
            text="Empty Sequencing Order Form",
            font=ctk.CTkFont(size=18, weight="bold"),
        ).pack(anchor="center")

        entry_frame_ef = ctk.CTkFrame(self.empty_form_frame, fg_color="transparent")
        entry_frame_ef.pack(fill="x", pady=(2, 2))

        self.entry_empty_form = ctk.CTkEntry(entry_frame_ef, border_width=1)
        self.entry_empty_form.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.button_empty_form = ctk.CTkButton(
            entry_frame_ef,
            text="Select",
            command=self.browse_empty_form,
            width=90,
        )
        self.button_empty_form.pack(side="right")

        self.empty_form_validation_label = ctk.CTkLabel(
            self.empty_form_frame,
            text="",
            font=ctk.CTkFont(size=14),
            anchor="w",
            justify="left",
        )
        self.empty_form_validation_label.pack(fill="x")

        # --- Sequencing Validation Template (Manual) ---
        self.seq_validation_file_frame = ctk.CTkFrame(master, fg_color="transparent")
        self.seq_validation_file_frame.pack(pady=15, padx=10, fill="x")

        ctk.CTkLabel(
            self.seq_validation_file_frame,
            text="Sequencing Validation Template",
            font=ctk.CTkFont(size=18, weight="bold"),
        ).pack(anchor="center")

        entry_frame_sv = ctk.CTkFrame(
            self.seq_validation_file_frame, fg_color="transparent"
        )
        entry_frame_sv.pack(fill="x", pady=(2, 2))

        self.entry_seq_validation_file = ctk.CTkEntry(entry_frame_sv, border_width=1)
        self.entry_seq_validation_file.pack(
            side="left", fill="x", expand=True, padx=(0, 10)
        )

        self.button_seq_validation_file = ctk.CTkButton(
            entry_frame_sv,
            text="Select",
            command=self.browse_seq_validation_file,
            width=90,
        )
        self.button_seq_validation_file.pack(side="right")

        self.seq_validation_validation_label = ctk.CTkLabel(
            self.seq_validation_file_frame,
            text="",
            font=ctk.CTkFont(size=14),
            anchor="w",
            justify="left",
        )
        self.seq_validation_validation_label.pack(fill="x")

        # --- Generate Button ---
        generate_button_frame = ctk.CTkFrame(master, fg_color="transparent")
        generate_button_frame.pack(pady=20, padx=20, fill="x")
        generate_button_frame.grid_columnconfigure((0, 2), weight=2)  # Spacers
        generate_button_frame.grid_columnconfigure(1, weight=1)  # Button column

        self.button_generate = ctk.CTkButton(
            generate_button_frame,
            text="Generate Form",
            command=self.generate_order_form,
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        )
        self.button_generate.configure(state="disabled", fg_color="gray50")
        self.button_generate.grid(row=0, column=1, sticky="ew")

        # --- Log Frame ---
        log_frame = ctk.CTkFrame(master)
        log_frame.pack(pady=(10, 15), padx=15, fill="both", expand=True)

        ctk.CTkLabel(
            log_frame,
            text="Log",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w", padx=10, pady=(5, 2))

        self.log_textbox = ctk.CTkTextbox(
            log_frame,
            activate_scrollbars=True,
            state="disabled",
            fg_color="white",
            text_color="black",
        )
        self.log_textbox.pack(expand=True, fill="both", padx=10, pady=(2, 10))

        # --- Setup initial paths and trigger validation ---
        # Defer validation to allow UI to render first
        self.master.after(
            200, lambda: self.setup_initial_paths(primer_log_path, empty_form_path)
        )

    def _log(self, message):
        if self.log_textbox and self.log_textbox.winfo_exists():
            self.log_textbox.insert("end", message + "\n")
            self.log_textbox.see("end")

    def _log_threaded(self, message):
        # Safely log from a worker thread
        if self.log_textbox:
            self.master.after(0, self._log, message)

    def _start_validation_thread(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        expected_sheet,
        validation_flag_name,
        next_callback=None,
    ):
        """Starts a background thread to validate a file path."""
        # Reset UI to neutral before starting validation
        self._update_validation_ui(
            entry_widget, validation_label_widget, None, "", validation_flag_name
        )

        thread = threading.Thread(
            target=self._validate_file,
            args=(
                entry_widget,
                validation_label_widget,
                file_path,
                expected_sheet,
                validation_flag_name,
                next_callback,
            ),
            daemon=True,
        )
        thread.start()

    def _validate_file(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        expected_sheet,
        validation_flag_name,
        next_callback=None,
    ):
        """
        Checks if a file exists and contains the expected sheet.
        This method runs in a background thread and schedules UI updates.
        """
        if not file_path:
            # This case is handled by the initial reset in _start_validation_thread
            return

        if not os.path.exists(file_path):
            message = "File not found."
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                message,
                validation_flag_name,
                next_callback,
            )
            return

        try:
            # For non-Excel files, this will raise an error quickly.
            wb = openpyxl.load_workbook(file_path, read_only=True)
            if expected_sheet and expected_sheet not in wb.sheetnames:
                message = f"Sheet '{expected_sheet}' not found."
                self.master.after(
                    0,
                    self._update_validation_ui,
                    entry_widget,
                    validation_label_widget,
                    False,
                    message,
                    validation_flag_name,
                    next_callback,
                )
                return

            # If we reach here, the file is valid.
            message = "Validated"
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                True,
                message,
                validation_flag_name,
                next_callback,
            )
        except Exception as e:
            # Catches errors from openpyxl (e.g., corrupted file, not an Excel file)
            message = "Not a valid Excel file."
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                message,
                validation_flag_name,
                next_callback,
            )

    def _update_validation_ui(
        self,
        entry_widget,
        validation_label_widget,
        is_valid,
        message,
        validation_flag_name,
        next_callback=None,
    ):
        """
        Updates the UI to reflect the validation status of a file.
        This method is always run on the main UI thread.
        """
        # Update the validation status flag for the specific file
        setattr(self, validation_flag_name, is_valid)

        if is_valid is True:
            entry_widget.configure(border_color="dark green", border_width=2)
            validation_label_widget.configure(text=f"✔ {message}", text_color="green")
        elif is_valid is False:
            entry_widget.configure(border_color="#D32F2F", border_width=2)
            validation_label_widget.configure(
                text=f"✖ Invalid: {message}", text_color="#D32F2F"
            )
        else:  # is_valid is None (resetting to neutral)
            entry_widget.configure(
                border_width=1,
                border_color=ctk.ThemeManager.theme["CTkEntry"]["border_color"],
            )
            validation_label_widget.configure(
                text="", text_color=ctk.ThemeManager.theme["CTkLabel"]["text_color"]
            )

        self._check_all_validations()

        if next_callback:
            next_callback()

    def _check_all_validations(self):
        """
        Checks if all files are valid and enables/disables the Generate button accordingly.
        """
        all_valid = (
            self.primerlog_valid and self.empty_form_valid and self.seq_validation_valid
        )
        if all_valid:
            self.button_generate.configure(state="normal", fg_color="#3C702D")
        else:
            self.button_generate.configure(state="disabled", fg_color="gray50")

    def setup_initial_paths(self, primer_log_path, empty_form_path):
        """
        Pre-fills entry widgets with provided paths and then starts sequential validation.
        """
        # --- Phase 1: Autofill all paths immediately ---
        if primer_log_path:
            self.entry_primerlog.delete(0, "end")
            self.entry_primerlog.insert(0, primer_log_path)

        if empty_form_path:
            self.entry_empty_form.delete(0, "end")
            self.entry_empty_form.insert(0, empty_form_path)

        # --- Phase 2: Start the validation chain ---
        self._start_sequential_validation()

    def _start_sequential_validation(self):
        """Starts the validation for any pre-filled paths in sequence."""
        primer_log_path = self.entry_primerlog.get()
        empty_form_path = self.entry_empty_form.get()

        # Define chained callbacks
        def validate_empty_form():
            if empty_form_path:
                self._start_validation_thread(
                    self.entry_empty_form,
                    self.empty_form_validation_label,
                    empty_form_path,
                    None,  # No specific sheet check needed
                    "empty_form_valid",
                )

        if primer_log_path:
            self._start_validation_thread(
                self.entry_primerlog,
                self.primerlog_validation_label,
                primer_log_path,
                "input addon",
                "primerlog_valid",
                next_callback=validate_empty_form,
            )
        else:
            validate_empty_form()  # Skip to next if path is missing

    def browse_primerlog(self):
        # Function to search for the primer log in the UI
        filename = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if filename:
            self.entry_primerlog.delete(0, "end")
            self.entry_primerlog.insert(0, filename)
            self._start_validation_thread(
                self.entry_primerlog,
                self.primerlog_validation_label,
                filename,
                "input addon",
                "primerlog_valid",
            )

    def browse_seq_validation_file(self):
        # Function to search for the Sequencing Validation file in the UI
        filename = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if filename:
            self.entry_seq_validation_file.delete(0, "end")
            self.entry_seq_validation_file.insert(0, filename)
            self._start_validation_thread(
                self.entry_seq_validation_file,
                self.seq_validation_validation_label,
                filename,
                "Sequencing Pending",
                "seq_validation_valid",
            )

    def browse_empty_form(self):
        # Function to search for the empty order form file (Template 3) in the UI
        filename = filedialog.askopenfilename(filetypes=[("Excel files", "*.xlsx")])
        if filename:
            self.entry_empty_form.delete(0, "end")
            self.entry_empty_form.insert(0, filename)
            self._start_validation_thread(
                self.entry_empty_form,
                self.empty_form_validation_label,
                filename,
                None,
                "empty_form_valid",
            )

    def generate_order_form(self):
        primerlog = self.entry_primerlog.get()
        seq_validation_file = self.entry_seq_validation_file.get()
        empty_form = self.entry_empty_form.get()

        if not seq_validation_file:
            messagebox.showerror(
                "Error", "Sequencing Validation Template path is missing."
            )
            return
        output_dir = os.path.dirname(seq_validation_file)

        # Disable button, create log window, and start worker thread
        self.button_generate.configure(state="disabled")
        self.log_textbox.configure(state="normal")
        self.log_textbox.delete("1.0", "end")
        self._log("Starting generation process...")

        thread = threading.Thread(
            target=self._generation_worker,
            args=(
                primerlog,
                seq_validation_file,
                empty_form,
                output_dir,
                self.seq_layout_path,
            ),
            daemon=True,
        )
        thread.start()

    def _update_links_in_layout(self, layout_path, new_source_path):
        """Updates external links in the copied layout Excel file using COM."""
        try:
            import pythoncom
            import win32com.client
        except ImportError:
            self._log_threaded("WARNING: pywin32 not installed. Skipping link update.")
            return

        self._log_threaded("Updating external links in layout file...")
        excel = None
        wb = None
        try:
            pythoncom.CoInitialize()
            excel = win32com.client.Dispatch("Excel.Application")
            excel.Visible = False
            excel.DisplayAlerts = False

            abs_path = os.path.abspath(layout_path)
            wb = excel.Workbooks.Open(abs_path)

            # xlExcelLinks = 1
            links = wb.LinkSources(1)
            updated = False

            if links:
                for link in links:
                    # Check if the link name contains the order form keyword
                    if "Sequencing_order_form" in os.path.basename(link):
                        self._log_threaded(f"  > Found link: {os.path.basename(link)}")
                        wb.ChangeLink(Name=link, NewName=new_source_path, Type=1)
                        self._log_threaded(
                            f"  > Updated to: {os.path.basename(new_source_path)}"
                        )
                        updated = True
                        break

            if updated:
                wb.Save()
                self._log_threaded("  > Layout file saved.")
            else:
                self._log_threaded("  > No matching links found.")

        except Exception as e:
            self._log_threaded(f"ERROR updating links: {e}")
        finally:
            if wb:
                try:
                    wb.Close(SaveChanges=False)
                except:
                    pass
            if excel:
                try:
                    excel.Quit()
                except:
                    pass
            pythoncom.CoUninitialize()

    def _generation_worker(
        self, primerlog, seq_validation_file, empty_form, output_dir, seq_layout_path
    ):
        try:
            self._log_threaded(
                f"Reading Sequencing Validation file: {os.path.basename(seq_validation_file)}"
            )
            try:
                wb = openpyxl.load_workbook(seq_validation_file, read_only=True)
                sequencing_sheet = wb["Sequencing Pending"]
            except Exception as e:
                self.master.after(
                    0,
                    lambda: messagebox.showerror(
                        "File Read Error",
                        f"Could not read the Sequencing Validation file: {e}",
                    ),
                )
                return

            self._log_threaded("Extracting unique Work No....")
            work_nos = []
            unique_work_nos = set()
            for row in sequencing_sheet.iter_rows(
                min_row=2, max_row=98, min_col=1, max_col=1, values_only=True
            ):
                if row[0] and row[0] not in unique_work_nos:
                    work_nos.append(row[0])
                    unique_work_nos.add(row[0])
            self._log_threaded(f"Found {len(work_nos)} unique Work No.")

            warnings, new_order_form_path = process_work_nos(
                work_nos,
                primerlog,
                seq_validation_file,
                empty_form,
                output_dir,
                logger=self._log_threaded,
            )

            status_messages = []
            has_warnings = False

            if warnings:
                has_warnings = True
                status_messages.append("Sequencing Order form generated with warnings:")
                status_messages.extend([f"• {w}" for w in warnings])
            else:
                status_messages.append("Sequencing Order form generated and saved!")
                self._log_threaded(
                    "\nSuccess! Sequencing Order form generated and saved."
                )

            if seq_layout_path and os.path.exists(seq_layout_path):
                try:
                    folder_name = os.path.basename(output_dir)
                    dest_filename = f"Sequencing_layout_{folder_name}.xlsx"
                    dest_path = os.path.join(output_dir, dest_filename)
                    shutil.copy2(seq_layout_path, dest_path)
                    self._log_threaded(
                        f"Successfully copied layout file: {dest_filename}"
                    )
                    self._update_links_in_layout(dest_path, new_order_form_path)
                    status_messages.append("Layout file copied and links updated.")
                except Exception as copy_e:
                    self._log_threaded(f"WARNING: Could not copy layout file: {copy_e}")
                    status_messages.append(f"WARNING: Layout file copy failed: {copy_e}")
                    has_warnings = True

            final_message = "\n".join(status_messages)
            if has_warnings:
                self.master.after(
                    0,
                    lambda: messagebox.showwarning("Generated with Warnings", final_message),
                )
            else:
                self.master.after(
                    0,
                    lambda: messagebox.showinfo("Success", final_message),
                )
        except Exception as e:
            self._log_threaded(f"\nERROR: {traceback.format_exc()}")
            self.master.after(0, lambda: messagebox.showerror("Error", str(e)))
        finally:
            self.master.after(0, self.button_generate.configure, {"state": "normal"})
            self._log_threaded("Process finished.")
            self.master.after(0, self.log_textbox.configure, {"state": "disabled"})


def main():
    # These variables act as placeholders for the paths that the App Hub
    # would pass to this application. When running standalone or if the
    # paths are not configured in the hub, they default to None.
    primer_log_path_from_hub = None
    empty_form_path_from_hub = None

    # Initialize the main CustomTkinter window.
    ctk.set_appearance_mode("System")  # Modes: "System" (default), "Dark", "Light"
    ctk.set_default_color_theme(
        "blue"
    )  # Themes: "blue" (default), "green", "dark-blue"

    root = ctk.CTk()
    # Create an instance of the GUI, passing the (potentially None) paths.
    # The GUI will handle autofilling the entry fields if the paths are valid.
    gui = OrderFormGUI(
        root,
        primer_log_path=primer_log_path_from_hub,
        empty_form_path=empty_form_path_from_hub,
    )
    # Start the event loop.
    root.mainloop()


if __name__ == "__main__":
    main()
