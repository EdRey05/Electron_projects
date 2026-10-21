"""
Bio Basic Inc. - Canada
Made by: Eduardo Reyes, Ph.D.
Contact: ed5reyes@outlook.com

Version: 4.2
Date: Oct 13, 2026
Notes: Fixed trailing dot bug in folder names on tab 3 output. Internal
    variables and log strings follow the canonical job-log naming
    (work_no).
"""

import os
import openpyxl
import shutil
import re
import threading
import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import datetime

import customtkinter as ctk


class SequencingFileOrganizerApp:
    """
    File organizer, reference distributor, and SG check document generator
    for the Gene Synthesis workflow.

    Three tabs, each with a pre-scan on Run:
    1. Organize Sequencing Files - Scan -> Abort/Skip/All -> Sort
    2. Distribute Reference Files - Scan -> Abort/Skip/All -> Copy
    3. SG Check Document - Scan -> Abort/Skip/All -> Generate
    """

    def __init__(
        self,
        master,
        reference_sequences_path=None,
        sorted_sequencing_files_path=None,
        sg_check_template_path=None,
        sg_check_output_path=None,
    ):
        self.master = master
        master.title("Sequencing File Organizer")
        master.geometry("850x750")

        # Hub-provided paths (passed via pass_template_as_kwarg and pass_dir_as_kwarg)
        self._hub_template_path = sg_check_template_path
        self._hub_sg_check_path = sg_check_output_path

        # --- Validation flags (path existence) ---
        self.source_dir_valid = False
        self.dest_dir_valid = False
        self.reference_dir_valid = False
        self.template_file_valid = False
        self.output_doc_dir_valid = False

        # --- Internal path storage ---
        self._source_dir_path = ""
        self._dest_dir_path = ""
        self._reference_dir_path = ""
        self._template_file_path = ""
        self._output_doc_dir_path = ""

        # --- Store processed results ---
        self.organized_work_nos = []
        self.created_folder_names = []

        # --- Manual folder list for Tab 3 ---
        self.manual_folders = []

        # --- Build UI ---
        self._setup_ui()

        # --- Setup initial paths ---
        self.master.after(
            200,
            lambda: self._setup_initial_paths(
                reference_sequences_path, sorted_sequencing_files_path
            ),
        )

    def _setup_ui(self):
        """Creates all UI widgets."""

        main_frame = ctk.CTkFrame(self.master, fg_color="transparent")
        main_frame.pack(fill="both", expand=True, padx=15, pady=15)

        self.tab_view = ctk.CTkTabview(main_frame)
        self.tab_view.pack(fill="both", expand=True)

        tab_1 = self.tab_view.add("Organize Files")
        tab_2 = self.tab_view.add("Distribute References")
        tab_3 = self.tab_view.add("SG Check Document")

        # ================================
        # TAB 1: Organize Sequencing Files
        # ================================
        tab1_frame = ctk.CTkFrame(tab_1, fg_color="transparent")
        tab1_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # Reference Files Directory
        ref_frame = ctk.CTkFrame(tab1_frame, fg_color="transparent")
        ref_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(
            ref_frame,
            text="Reference Sequences Directory",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        ref_entry_frame = ctk.CTkFrame(ref_frame, fg_color="transparent")
        ref_entry_frame.pack(fill="x", pady=(5, 0))

        self.entry_reference_dir = ctk.CTkEntry(ref_entry_frame, border_width=1)
        self.entry_reference_dir.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(
            ref_entry_frame, text="Browse", command=self._browse_reference_dir, width=80
        ).pack(side="right")

        self.reference_validation_label = ctk.CTkLabel(
            ref_frame, text="", font=ctk.CTkFont(size=13), anchor="w", justify="left"
        )
        self.reference_validation_label.pack(fill="x", pady=(2, 0))

        # Destination Directory
        dest_frame = ctk.CTkFrame(tab1_frame, fg_color="transparent")
        dest_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(
            dest_frame,
            text="Sorted Sequencing Folder",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        dest_entry_frame = ctk.CTkFrame(dest_frame, fg_color="transparent")
        dest_entry_frame.pack(fill="x", pady=(5, 0))

        self.entry_dest_dir = ctk.CTkEntry(dest_entry_frame, border_width=1)
        self.entry_dest_dir.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(
            dest_entry_frame, text="Browse", command=self._browse_dest_dir, width=80
        ).pack(side="right")

        self.dest_validation_label = ctk.CTkLabel(
            dest_frame, text="", font=ctk.CTkFont(size=13), anchor="w", justify="left"
        )
        self.dest_validation_label.pack(fill="x", pady=(2, 0))

        # Source Directory
        source_frame = ctk.CTkFrame(tab1_frame, fg_color="transparent")
        source_frame.pack(fill="x", pady=(0, 15))

        ctk.CTkLabel(
            source_frame,
            text="Sequencing Files Directory",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        source_entry_frame = ctk.CTkFrame(source_frame, fg_color="transparent")
        source_entry_frame.pack(fill="x", pady=(5, 0))

        self.entry_source_dir = ctk.CTkEntry(source_entry_frame, border_width=1)
        self.entry_source_dir.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(
            source_entry_frame, text="Browse", command=self._browse_source_dir, width=80
        ).pack(side="right")

        self.source_validation_label = ctk.CTkLabel(
            source_frame, text="", font=ctk.CTkFont(size=13), anchor="w", justify="left"
        )
        self.source_validation_label.pack(fill="x", pady=(2, 0))

        # Organize Button
        self.btn_organize = ctk.CTkButton(
            tab1_frame,
            text="Organize Files",
            command=self._start_organize,
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        )
        self.btn_organize.configure(state="disabled", fg_color="gray50")
        self.btn_organize.pack(pady=15)

        self.status_label_1 = ctk.CTkLabel(
            tab1_frame, text="", font=ctk.CTkFont(size=13), text_color="gray"
        )
        self.status_label_1.pack(pady=5)

        # ================================
        # TAB 2: Distribute Reference Files
        # ================================
        tab2_frame = ctk.CTkFrame(tab_2, fg_color="transparent")
        tab2_frame.pack(fill="both", expand=True, padx=10, pady=10)

        info_frame = ctk.CTkFrame(tab2_frame, fg_color="transparent")
        info_frame.pack(fill="x", pady=(0, 15))

        ctk.CTkLabel(
            info_frame,
            text="Destination Directory",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")
        ctk.CTkLabel(
            info_frame,
            text="(Uses the 'Sorted Sequencing Folder' from Tab 1)",
            font=ctk.CTkFont(size=12),
            text_color="gray",
        ).pack(anchor="w", pady=(2, 0))

        # Distribute Button
        self.btn_distribute = ctk.CTkButton(
            tab2_frame,
            text="Distribute References",
            command=self._start_distribute,
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        )
        self.btn_distribute.configure(state="disabled", fg_color="gray50")
        self.btn_distribute.pack(pady=15)

        self.status_label_2 = ctk.CTkLabel(
            tab2_frame, text="", font=ctk.CTkFont(size=13), text_color="gray"
        )
        self.status_label_2.pack(pady=5)

        # ================================
        # TAB 3: SG Check Document
        # ================================
        tab3_frame = ctk.CTkFrame(tab_3, fg_color="transparent")
        tab3_frame.pack(fill="both", expand=True, padx=10, pady=10)

        # Template Selection
        temp_frame = ctk.CTkFrame(tab3_frame, fg_color="transparent")
        temp_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(
            temp_frame,
            text="Excel Template File",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        temp_entry_frame = ctk.CTkFrame(temp_frame, fg_color="transparent")
        temp_entry_frame.pack(fill="x", pady=(5, 0))

        self.entry_template_file = ctk.CTkEntry(temp_entry_frame, border_width=1)
        self.entry_template_file.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(
            temp_entry_frame,
            text="Browse",
            command=self._browse_template_file,
            width=80,
        ).pack(side="right")

        self.template_validation_label = ctk.CTkLabel(
            temp_frame, text="", font=ctk.CTkFont(size=13), anchor="w", justify="left"
        )
        self.template_validation_label.pack(fill="x", pady=(2, 0))

        # Output Directory
        out_doc_frame = ctk.CTkFrame(tab3_frame, fg_color="transparent")
        out_doc_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(
            out_doc_frame,
            text="Document Destination Folder",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        out_doc_entry_frame = ctk.CTkFrame(out_doc_frame, fg_color="transparent")
        out_doc_entry_frame.pack(fill="x", pady=(5, 0))

        self.entry_output_doc_dir = ctk.CTkEntry(out_doc_entry_frame, border_width=1)
        self.entry_output_doc_dir.pack(side="left", fill="x", expand=True, padx=(0, 10))
        ctk.CTkButton(
            out_doc_entry_frame,
            text="Browse",
            command=self._browse_output_doc_dir,
            width=80,
        ).pack(side="right")

        self.output_doc_validation_label = ctk.CTkLabel(
            out_doc_frame,
            text="",
            font=ctk.CTkFont(size=13),
            anchor="w",
            justify="left",
        )
        self.output_doc_validation_label.pack(fill="x", pady=(2, 0))

        # Folder source selection
        folder_source_frame = ctk.CTkFrame(tab3_frame, fg_color="transparent")
        folder_source_frame.pack(fill="x", pady=(0, 10))

        ctk.CTkLabel(
            folder_source_frame,
            text="Folder Names Source",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w")

        self.check_use_organized = ctk.CTkCheckBox(
            folder_source_frame,
            text="Use organized folders from Tab 1",
            command=self._on_folder_source_change,
        )
        self.check_use_organized.pack(anchor="w", pady=(5, 2))

        self.organized_info_label = ctk.CTkLabel(
            folder_source_frame,
            text="No folders organized yet. Run Tab 1 first.",
            font=ctk.CTkFont(size=12),
            text_color="gray",
        )
        self.organized_info_label.pack(anchor="w", pady=(0, 5))

        # Manual folder addition
        add_entry_frame = ctk.CTkFrame(folder_source_frame, fg_color="transparent")
        add_entry_frame.pack(fill="x", pady=(0, 5))

        self.entry_add_folder = ctk.CTkEntry(
            add_entry_frame,
            border_width=1,
            placeholder_text="Type folder name or drop a folder path here...",
        )
        self.entry_add_folder.pack(side="left", fill="x", expand=True, padx=(0, 5))

        ctk.CTkButton(
            add_entry_frame,
            text="Add",
            command=self._add_manual_folder,
            width=60,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        ).pack(side="right")
        ctk.CTkButton(
            add_entry_frame, text="Browse", command=self._browse_add_folder, width=70
        ).pack(side="right", padx=(0, 5))

        self.manual_folders_frame = ctk.CTkScrollableFrame(
            folder_source_frame, height=100, fg_color="transparent"
        )
        self.manual_folders_frame.pack(fill="x", pady=(0, 5))

        ctk.CTkLabel(
            self.manual_folders_frame,
            text="No folders added manually.",
            font=ctk.CTkFont(size=12),
            text_color="gray",
        ).pack(anchor="w")

        # Generate Button
        self.btn_generate_sg = ctk.CTkButton(
            tab3_frame,
            text="Generate SG Check Document",
            command=self._start_generate_sg_check,
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        )
        self.btn_generate_sg.configure(state="disabled", fg_color="gray50")
        self.btn_generate_sg.pack(pady=10)

        self.status_label_3 = ctk.CTkLabel(
            tab3_frame, text="", font=ctk.CTkFont(size=13), text_color="gray"
        )
        self.status_label_3.pack(pady=5)

        # ================================
        # LOG FRAME
        # ================================
        log_frame = ctk.CTkFrame(main_frame)
        log_frame.pack(fill="both", expand=True, pady=(15, 0))

        ctk.CTkLabel(
            log_frame, text="Log", font=ctk.CTkFont(size=14, weight="bold")
        ).pack(anchor="w", padx=10, pady=(5, 2))

        self.log_textbox = ctk.CTkTextbox(
            log_frame,
            activate_scrollbars=True,
            state="disabled",
            fg_color="white",
            text_color="black",
        )
        self.log_textbox.pack(expand=True, fill="both", padx=10, pady=(2, 10))

    # ─── Logging ────────────────────────────────────────────────

    def _log(self, message):
        if self.log_textbox and self.log_textbox.winfo_exists():
            self.log_textbox.configure(state="normal")
            self.log_textbox.insert("end", message + "\n")
            self.log_textbox.see("end")
            self.log_textbox.configure(state="disabled")

    def _log_threaded(self, message):
        if self.log_textbox:
            self.master.after(0, self._log, message)

    def _clear_log(self):
        if self.log_textbox and self.log_textbox.winfo_exists():
            self.log_textbox.configure(state="normal")
            self.log_textbox.delete("1.0", "end")
            self.log_textbox.configure(state="disabled")

    # ─── Path Validation ────────────────────────────────────────

    def _start_validation_thread(
        self,
        entry_widget,
        validation_label_widget,
        dir_path,
        validation_flag_name,
        dir_display_name,
        next_callback=None,
    ):
        self._update_validation_ui(
            entry_widget, validation_label_widget, None, "", validation_flag_name
        )
        thread = threading.Thread(
            target=self._validate_directory,
            args=(
                entry_widget,
                validation_label_widget,
                dir_path,
                validation_flag_name,
                dir_display_name,
                next_callback,
            ),
            daemon=True,
        )
        thread.start()

    def _validate_directory(
        self,
        entry_widget,
        validation_label_widget,
        dir_path,
        validation_flag_name,
        dir_display_name,
        next_callback=None,
    ):
        if not dir_path:
            return
        if not os.path.exists(dir_path) or not os.path.isdir(dir_path):
            msg = (
                "Directory not found."
                if not os.path.exists(dir_path)
                else "Path is not a directory."
            )
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                msg,
                validation_flag_name,
                dir_display_name,
                next_callback,
            )
            return
        self.master.after(
            0,
            self._update_validation_ui,
            entry_widget,
            validation_label_widget,
            True,
            "Validated",
            validation_flag_name,
            dir_display_name,
            next_callback,
        )

    def _validate_file(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        validation_flag_name,
        file_display_name,
        next_callback=None,
    ):
        if not file_path:
            return
        if not os.path.exists(file_path) or not os.path.isfile(file_path):
            msg = (
                "File not found."
                if not os.path.exists(file_path)
                else "Path is not a file."
            )
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                msg,
                validation_flag_name,
                file_display_name,
                next_callback,
            )
            return
        self.master.after(
            0,
            self._update_validation_ui,
            entry_widget,
            validation_label_widget,
            True,
            "Validated",
            validation_flag_name,
            file_display_name,
            next_callback,
        )

    def _start_file_validation_thread(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        validation_flag_name,
        file_display_name,
    ):
        self._update_validation_ui(
            entry_widget, validation_label_widget, None, "", validation_flag_name
        )
        thread = threading.Thread(
            target=self._validate_file,
            args=(
                entry_widget,
                validation_label_widget,
                file_path,
                validation_flag_name,
                file_display_name,
            ),
            daemon=True,
        )
        thread.start()

    def _update_validation_ui(
        self,
        entry_widget,
        validation_label_widget,
        is_valid,
        message,
        validation_flag_name,
        dir_display_name=None,
        next_callback=None,
    ):
        setattr(self, validation_flag_name, is_valid)

        if is_valid is True:
            if validation_flag_name == "source_dir_valid":
                self._source_dir_path = entry_widget.get()
            elif validation_flag_name == "dest_dir_valid":
                self._dest_dir_path = entry_widget.get()
            elif validation_flag_name == "reference_dir_valid":
                self._reference_dir_path = entry_widget.get()
            elif validation_flag_name == "template_file_valid":
                self._template_file_path = entry_widget.get()
            elif validation_flag_name == "output_doc_dir_valid":
                self._output_doc_dir_path = entry_widget.get()
            entry_widget.configure(border_color="dark green", border_width=2)
            validation_label_widget.configure(text=f"✔ {message}", text_color="green")
        elif is_valid is False:
            entry_widget.configure(border_color="#D32F2F", border_width=2)
            validation_label_widget.configure(
                text=f"✖ Invalid: {message}", text_color="#D32F2F"
            )
        else:
            entry_widget.configure(
                border_width=1,
                border_color=ctk.ThemeManager.theme["CTkEntry"]["border_color"],
            )
            validation_label_widget.configure(
                text="", text_color=ctk.ThemeManager.theme["CTkLabel"]["text_color"]
            )

        self._check_all_validations()
        if next_callback:
            next_callback()

    def _check_all_validations(self):
        """Enables Run buttons when paths are valid and conditions are met."""
        if self.source_dir_valid and self.dest_dir_valid:
            self.btn_organize.configure(state="normal", fg_color="#3C702D")
        else:
            self.btn_organize.configure(state="disabled", fg_color="gray50")

        if self.reference_dir_valid and self.organized_work_nos:
            self.btn_distribute.configure(state="normal", fg_color="#3C702D")
        else:
            self.btn_distribute.configure(state="disabled", fg_color="gray50")

        has_folders = len(self._get_combined_folders()) > 0
        if self.template_file_valid and self.output_doc_dir_valid and has_folders:
            self.btn_generate_sg.configure(state="normal", fg_color="#3C702D")
        else:
            self.btn_generate_sg.configure(state="disabled", fg_color="gray50")

    # ─── Pre-execution Scan Dialog ──────────────────────────────

    def _show_issues_dialog(self, title, issues_list, summary_text):
        """
        Show a dialog with issues found and 3 options.
        Returns: 'abort', 'skip', or 'all'
        """
        dialog = tk.Toplevel(self.master)
        dialog.title(title)
        dialog.geometry("540x380")
        dialog.transient(self.master)
        dialog.grab_set()
        dialog.resizable(True, True)

        result = {"action": "abort"}

        tk.Label(
            dialog,
            text=summary_text,
            font=("Arial", 12, "bold"),
            wraplength=500,
            justify="left",
        ).pack(padx=15, pady=(15, 5), anchor="w")

        # Scrollable issues list
        text_frame = tk.Frame(dialog)
        text_frame.pack(fill="both", expand=True, padx=15, pady=5)

        scrollbar = tk.Scrollbar(text_frame)
        scrollbar.pack(side="right", fill="y")

        issues_textbox = tk.Text(
            text_frame,
            wrap="word",
            font=("Consolas", 10),
            yscrollcommand=scrollbar.set,
            height=10,
        )
        for line in issues_list:
            issues_textbox.insert("end", line + "\n")
        issues_textbox.configure(state="disabled")
        issues_textbox.pack(fill="both", expand=True)
        scrollbar.configure(command=issues_textbox.yview)

        tk.Label(dialog, text="Choose how to proceed:", font=("Arial", 11)).pack(
            pady=(5, 5)
        )

        btn_frame = tk.Frame(dialog)
        btn_frame.pack(pady=(0, 15))

        tk.Button(
            btn_frame,
            text="Abort",
            width=10,
            bg="#D32F2F",
            fg="white",
            font=("Arial", 11, "bold"),
            activebackground="#B71C1C",
            command=lambda: (result.__setitem__("action", "abort"), dialog.destroy()),
        ).pack(side="left", padx=5)
        tk.Button(
            btn_frame,
            text="Skip Issues",
            width=10,
            bg="#eab308",
            fg="black",
            font=("Arial", 11, "bold"),
            activebackground="#ca8a04",
            command=lambda: (result.__setitem__("action", "skip"), dialog.destroy()),
        ).pack(side="left", padx=5)
        tk.Button(
            btn_frame,
            text="Process All",
            width=12,
            bg="#3C702D",
            fg="white",
            font=("Arial", 11, "bold"),
            activebackground="#2d5a22",
            command=lambda: (result.__setitem__("action", "all"), dialog.destroy()),
        ).pack(side="left", padx=5)

        dialog.wait_window()
        return result["action"]

    # ─── Manual Folder Management ───────────────────────────────

    def _add_manual_folder(self):
        folder = self.entry_add_folder.get().strip()
        if not folder:
            return
        folder_name = os.path.basename(folder.rstrip(os.sep))
        if folder_name and folder_name not in self.manual_folders:
            self.manual_folders.append(folder_name)
            self._refresh_manual_folders_list()
            self.entry_add_folder.delete(0, "end")
            self._check_all_validations()

    def _browse_add_folder(self):
        folder = filedialog.askdirectory(title="Select Folder to Add")
        if folder:
            folder_name = os.path.basename(folder)
            if folder_name not in self.manual_folders:
                self.manual_folders.append(folder_name)
                self._refresh_manual_folders_list()
                self._check_all_validations()

    def _remove_manual_folder(self, folder_name):
        if folder_name in self.manual_folders:
            self.manual_folders.remove(folder_name)
            self._refresh_manual_folders_list()
            self._check_all_validations()

    def _refresh_manual_folders_list(self):
        for widget in self.manual_folders_frame.winfo_children():
            widget.destroy()
        if not self.manual_folders:
            ctk.CTkLabel(
                self.manual_folders_frame,
                text="No folders added manually.",
                font=ctk.CTkFont(size=12),
                text_color="gray",
            ).pack(anchor="w")
            return
        for folder_name in self.manual_folders:
            row = ctk.CTkFrame(self.manual_folders_frame, fg_color="transparent")
            row.pack(fill="x", pady=1)
            ctk.CTkLabel(row, text=f"📁 {folder_name}", font=ctk.CTkFont(size=12)).pack(
                side="left"
            )
            ctk.CTkButton(
                row,
                text="✕",
                width=28,
                height=28,
                fg_color="#D32F2F",
                hover_color="#B71C1C",
                font=ctk.CTkFont(size=12),
                command=lambda fn=folder_name: self._remove_manual_folder(fn),
            ).pack(side="right")

    def _on_folder_source_change(self):
        self._check_all_validations()

    def _get_combined_folders(self):
        folders = []
        if self.check_use_organized.get():
            folders.extend(self.created_folder_names)
        folders.extend(self.manual_folders)
        seen = set()
        result = []
        for f in folders:
            if f not in seen:
                seen.add(f)
                result.append(f)
        return result

    # ─── Initial Path Setup ─────────────────────────────────────

    def _setup_initial_paths(
        self, reference_sequences_path, sorted_sequencing_files_path
    ):
        if reference_sequences_path:
            self.entry_reference_dir.delete(0, "end")
            self.entry_reference_dir.insert(0, reference_sequences_path)
        if sorted_sequencing_files_path:
            self.entry_dest_dir.delete(0, "end")
            self.entry_dest_dir.insert(0, sorted_sequencing_files_path)
        if self._hub_template_path and os.path.exists(self._hub_template_path):
            self.entry_template_file.delete(0, "end")
            self.entry_template_file.insert(0, self._hub_template_path)
        if self._hub_sg_check_path and os.path.exists(self._hub_sg_check_path):
            self.entry_output_doc_dir.delete(0, "end")
            self.entry_output_doc_dir.insert(0, self._hub_sg_check_path)
        self._start_sequential_validation()

    def _start_sequential_validation(self):
        source_dir = self.entry_source_dir.get()
        dest_dir = self.entry_dest_dir.get()
        reference_dir = self.entry_reference_dir.get()
        template_file = self.entry_template_file.get()
        output_doc_dir = self.entry_output_doc_dir.get()

        def validate_output_doc():
            if output_doc_dir:
                self._start_validation_thread(
                    self.entry_output_doc_dir,
                    self.output_doc_validation_label,
                    output_doc_dir,
                    "output_doc_dir_valid",
                    "Document Destination Folder",
                )

        def validate_template():
            if template_file:
                self._start_file_validation_thread(
                    self.entry_template_file,
                    self.template_validation_label,
                    template_file,
                    "template_file_valid",
                    "Excel Template File",
                )
            validate_output_doc()

        def validate_reference():
            if reference_dir:
                self._start_validation_thread(
                    self.entry_reference_dir,
                    self.reference_validation_label,
                    reference_dir,
                    "reference_dir_valid",
                    "Reference Files Directory",
                )
            validate_template()

        def validate_dest():
            if dest_dir:
                self._start_validation_thread(
                    self.entry_dest_dir,
                    self.dest_validation_label,
                    dest_dir,
                    "dest_dir_valid",
                    "Sorted Sequencing Folder",
                    next_callback=validate_reference,
                )
            else:
                validate_reference()

        if source_dir:
            self._start_validation_thread(
                self.entry_source_dir,
                self.source_validation_label,
                source_dir,
                "source_dir_valid",
                "Sequencing Files Directory",
                next_callback=validate_dest,
            )
        else:
            validate_dest()

    # ─── Browse Methods ─────────────────────────────────────────

    def _browse_source_dir(self):
        d = filedialog.askdirectory(title="Select Sequencing Files Directory")
        if d:
            self.entry_source_dir.delete(0, "end")
            self.entry_source_dir.insert(0, d)
            self._start_validation_thread(
                self.entry_source_dir,
                self.source_validation_label,
                d,
                "source_dir_valid",
                "Sequencing Files Directory",
            )

    def _browse_dest_dir(self):
        d = filedialog.askdirectory(title="Select Sorted Sequencing Folder")
        if d:
            self.entry_dest_dir.delete(0, "end")
            self.entry_dest_dir.insert(0, d)
            self._start_validation_thread(
                self.entry_dest_dir,
                self.dest_validation_label,
                d,
                "dest_dir_valid",
                "Sorted Sequencing Folder",
            )

    def _browse_reference_dir(self):
        d = filedialog.askdirectory(title="Select Reference Files Directory")
        if d:
            self.entry_reference_dir.delete(0, "end")
            self.entry_reference_dir.insert(0, d)
            self._start_validation_thread(
                self.entry_reference_dir,
                self.reference_validation_label,
                d,
                "reference_dir_valid",
                "Reference Files Directory",
            )

    def _browse_template_file(self):
        f = filedialog.askopenfilename(
            title="Select Excel Template", filetypes=[("Excel files", "*.xlsx *.xlsm")]
        )
        if f:
            self.entry_template_file.delete(0, "end")
            self.entry_template_file.insert(0, f)
            self._start_file_validation_thread(
                self.entry_template_file,
                self.template_validation_label,
                f,
                "template_file_valid",
                "Excel Template File",
            )

    def _browse_output_doc_dir(self):
        d = filedialog.askdirectory(title="Select Destination Folder for Document")
        if d:
            self.entry_output_doc_dir.delete(0, "end")
            self.entry_output_doc_dir.insert(0, d)
            self._start_validation_thread(
                self.entry_output_doc_dir,
                self.output_doc_validation_label,
                d,
                "output_doc_dir_valid",
                "Document Destination Folder",
            )

    # ─── Tab 1: Organize Files ───────────────────────────────────

    def _start_organize(self):
        """Scan source files, show issues dialog if needed, then organize."""
        source_dir = self._source_dir_path
        dest_dir = self._dest_dir_path

        # Pre-scan
        unparseable = []
        parseable = []
        seq_count = 0
        pos_count = 0

        for filename in os.listdir(source_dir):
            file_path = os.path.join(source_dir, filename)
            if not os.path.isfile(file_path):
                continue
            if filename.endswith(".seq"):
                seq_count += 1
                continue
            if (
                filename.startswith("POS")
                or filename.startswith("Pos")
                or filename.startswith("pos")
            ):
                pos_count += 1
                continue
            if filename.endswith(".ab1") or filename.endswith(".fasta"):
                work_no, plasmid_no = self._extract_info(filename)
                if work_no and plasmid_no:
                    folder_name = f"{work_no}.{plasmid_no}".replace("..", ".").rstrip(".")
                    parseable.append((filename, work_no, plasmid_no, folder_name))
                else:
                    unparseable.append(filename)

        # Check existing folders
        existing_folders = sorted(
            set(
                fn
                for _, _, _, fn in parseable
                if os.path.exists(os.path.join(dest_dir, fn))
            )
        )

        issues = []
        if unparseable:
            issues.append(f"⚠ {len(unparseable)} file(s) cannot be parsed:")
            for fn in unparseable:
                issues.append(f"  ✖ {fn}")
            issues.append("")

        if existing_folders:
            issues.append(f"📁 {len(existing_folders)} folder(s) already exist:")
            for fn in existing_folders[:10]:
                issues.append(f"  📁 {fn}")
            if len(existing_folders) > 10:
                issues.append(f"  ... and {len(existing_folders) - 10} more")
            issues.append("")

        if issues:
            summary = (
                f"Scan found {len(parseable)} parseable files"
                f"{f', {len(unparseable)} unparseable' if unparseable else ''}"
                f"{f', {len(existing_folders)} existing folders' if existing_folders else ''}."
            )

            action = self._show_issues_dialog(
                "Organize Files — Issues Found", issues, summary
            )

            if action == "abort":
                return
            # 'skip' = unparseable files are already not in parseable list
            # 'all' = same as skip for Tab 1 (unparseable can't be organized anyway)

        # Run
        self.btn_organize.configure(state="disabled", fg_color="gray50")
        self._clear_log()
        self._log("Starting file organization...")
        thread = threading.Thread(
            target=self._organize_thread, args=(parseable,), daemon=True
        )
        thread.start()

    def _organize_thread(self, parseable_files):
        try:
            source_dir = self._source_dir_path
            dest_dir = self._dest_dir_path

            self._log_threaded(f"Source: {source_dir}")
            self._log_threaded(f"Destination: {dest_dir}")

            # Delete .seq files
            self._log_threaded("\nDeleting .seq files from source...")
            deleted_count = 0
            for filename in os.listdir(source_dir):
                file_path = os.path.join(source_dir, filename)
                if os.path.isfile(file_path) and filename.endswith(".seq"):
                    try:
                        os.remove(file_path)
                        deleted_count += 1
                    except Exception as e:
                        self._log_threaded(f"  Error deleting {filename}: {e}")
            self._log_threaded(f"  Deleted {deleted_count} .seq files")

            # Log skipped POS control files
            pos_count = sum(
                1
                for fn in os.listdir(source_dir)
                if os.path.isfile(os.path.join(source_dir, fn))
                and (
                    fn.startswith("POS") or fn.startswith("Pos") or fn.startswith("pos")
                )
            )
            if pos_count:
                self._log_threaded(
                    f"\nSkipped {pos_count} positive control file(s) (POS prefix)."
                )

            # Organize
            self._log_threaded("\nOrganizing files...")
            organized_work_nos = []
            created_folder_names = []
            organized_count = 0

            for filename, work_no, plasmid_no, folder_name in parseable_files:
                file_path = os.path.join(source_dir, filename)
                if not os.path.exists(file_path):
                    self._log_threaded(f"  ⚠ Skipped (already moved?): {filename}")
                    continue

                folder_path = os.path.join(dest_dir, folder_name)
                if not os.path.exists(folder_path):
                    os.makedirs(folder_path)
                    self._log_threaded(f"  Created folder: {folder_name}")

                destination_file = os.path.join(folder_path, filename)
                try:
                    shutil.move(file_path, destination_file)
                    organized_work_nos.append(work_no)
                    if folder_name not in created_folder_names:
                        created_folder_names.append(folder_name)
                    organized_count += 1
                except Exception as e:
                    self._log_threaded(f"  Error moving {filename}: {e}")

            self.organized_work_nos = list(set(organized_work_nos))
            self.created_folder_names = sorted(created_folder_names)
            self._log_threaded(f"\nOrganized {organized_count} files.")
            self._log_threaded(f"Found {len(self.organized_work_nos)} unique Work No.")

            self.master.after(
                0,
                lambda: self.organized_info_label.configure(
                    text=f"{len(self.created_folder_names)} folders organized and ready for SG Check."
                ),
            )
            self.master.after(0, lambda: self.check_use_organized.select())
            self.master.after(0, self._check_all_validations)
            self.master.after(
                0,
                lambda: self.status_label_1.configure(
                    text=f"Organized {organized_count} files."
                ),
            )
            self._log_threaded("\n--- Organization Complete ---")

        except Exception as e:
            self._log_threaded(f"\nERROR: {e}")
            self.master.after(0, lambda: messagebox.showerror("Error", str(e)))
        finally:
            self.master.after(
                0,
                lambda: self.btn_organize.configure(state="normal", fg_color="#3C702D"),
            )

    # ─── Tab 2: Distribute References ───────────────────────────

    def _start_distribute(self):
        """Scan reference matches, show issues dialog, then distribute."""
        source_dir = self._reference_dir_path
        dest_base_dir = self._dest_dir_path

        # Pre-scan
        file_map = {}
        for filename in os.listdir(source_dir):
            if filename.endswith(".txt"):
                base_name = filename.split("+")[0].split(".")[0]
                if base_name in self.organized_work_nos:
                    if base_name not in file_map:
                        file_map[base_name] = []
                    file_map[base_name].append(filename)

        missing = sorted(
            set(work_no for work_no in self.organized_work_nos if work_no not in file_map)
        )

        # Check destination folders
        missing_folders = []
        for work_no in self.organized_work_nos:
            found = False
            for folder in os.listdir(dest_base_dir):
                if folder.split(".")[0] == work_no:
                    found = True
                    break
            if not found:
                missing_folders.append(work_no)

        issues = []
        if missing:
            issues.append(f"⚠ {len(missing)} Work No. missing reference .txt files:")
            for work_no in missing:
                issues.append(f"  ✖ {work_no}")
            issues.append("")

        if missing_folders:
            issues.append(
                f"⚠ {len(missing_folders)} Work No. have no destination folder:"
            )
            for work_no in sorted(set(missing_folders)):
                issues.append(f"  ✖ {work_no}")
            issues.append("")

        work_nos_to_skip = set()

        if issues:
            summary = (
                f"Scan found .txt files for {len(file_map)} of {len(self.organized_work_nos)} Work No."
                f"{f', {len(missing)} missing references' if missing else ''}"
                f"{f', {len(missing_folders)} missing folders' if missing_folders else ''}."
            )

            action = self._show_issues_dialog(
                "Distribute References — Issues Found", issues, summary
            )

            if action == "abort":
                return
            elif action == "skip":
                work_nos_to_skip = set(missing) | set(missing_folders)

        # Run
        self.btn_distribute.configure(state="disabled", fg_color="gray50")
        self._log("\nStarting reference distribution...")
        thread = threading.Thread(
            target=self._distribute_thread,
            args=(file_map, work_nos_to_skip),
            daemon=True,
        )
        thread.start()

    def _distribute_thread(self, file_map, work_nos_to_skip):
        try:
            source_dir = self._reference_dir_path
            dest_base_dir = self._dest_dir_path

            self._log_threaded(f"Reference files: {source_dir}")
            self._log_threaded(f"Destination: {dest_base_dir}")

            if work_nos_to_skip:
                self._log_threaded(
                    f"\nSkipping {len(work_nos_to_skip)} Work No. with issues."
                )

            self._log_threaded(f"\nDistributing reference files...")
            distributed_count = 0

            for folder in os.listdir(dest_base_dir):
                folder_base_name = folder.split(".")[0]
                if (
                    folder_base_name in self.organized_work_nos
                    and folder_base_name in file_map
                ):
                    if folder_base_name in work_nos_to_skip:
                        self._log_threaded(f"  Skipped {folder} (missing files)")
                        continue
                    destination_dir = os.path.join(dest_base_dir, folder)
                    for source_filename in file_map[folder_base_name]:
                        source_file = os.path.join(source_dir, source_filename)
                        destination_file = os.path.join(
                            destination_dir, source_filename
                        )
                        try:
                            shutil.copy(source_file, destination_file)
                            self._log_threaded(
                                f"  Copied {source_filename} to {folder}"
                            )
                            distributed_count += 1
                        except Exception as e:
                            self._log_threaded(
                                f"  Error copying {source_filename}: {e}"
                            )

            self._log_threaded(f"\nDistributed {distributed_count} reference files.")
            self.master.after(
                0,
                lambda: self.status_label_2.configure(
                    text=f"Distributed {distributed_count} files."
                ),
            )
            self._log_threaded("\n--- Distribution Complete ---")

        except Exception as e:
            self._log_threaded(f"\nERROR: {e}")
            self.master.after(0, lambda: messagebox.showerror("Error", str(e)))
        finally:
            self.master.after(
                0,
                lambda: self.btn_distribute.configure(
                    state="normal", fg_color="#3C702D"
                ),
            )

    # ─── Tab 3: SG Check Document ───────────────────────────────

    def _start_generate_sg_check(self):
        """Scan folder contents, show issues dialog, then generate."""
        folder_names = self._get_combined_folders()
        dest_dir = self._dest_dir_path

        # Pre-scan
        invalid_folders = []
        valid_folders = []

        for folder_name in folder_names:
            folder_path = os.path.join(dest_dir, folder_name)
            if not os.path.exists(folder_path) or not os.path.isdir(folder_path):
                invalid_folders.append((folder_name, "folder not found"))
                continue

            has_ab1 = any(
                fn.endswith(".ab1") or fn.endswith(".fasta")
                for fn in os.listdir(folder_path)
            )
            has_txt = any(fn.endswith(".txt") for fn in os.listdir(folder_path))

            missing_types = []
            if not has_ab1:
                missing_types.append("no .ab1/.fasta")
            if not has_txt:
                missing_types.append("no .txt")

            if missing_types:
                invalid_folders.append((folder_name, ", ".join(missing_types)))
            else:
                valid_folders.append(folder_name)

        issues = []
        if invalid_folders:
            issues.append(f"⚠ {len(invalid_folders)} folder(s) have missing files:")
            for name, issue in invalid_folders:
                issues.append(f"  ✖ {name}: {issue}")
            issues.append("")

        if issues:
            summary = (
                f"Scan found {len(valid_folders)} valid, {len(invalid_folders)} incomplete folder(s).\n"
                f"'Skip' = exclude incomplete folders from the Excel\n"
                f"'Process All' = include all folders (remove extras manually later)"
            )

            action = self._show_issues_dialog(
                "SG Check Document — Issues Found", issues, summary
            )

            if action == "abort":
                return
            elif action == "skip":
                folder_names = valid_folders

        # Run
        self.btn_generate_sg.configure(state="disabled", fg_color="gray50")
        self._log("\nGenerating SG Check Document...")
        thread = threading.Thread(
            target=self._generate_sg_check_thread, args=(folder_names,), daemon=True
        )
        thread.start()

    def _generate_sg_check_thread(self, folder_names):
        try:
            template_path = self._template_file_path
            output_dir = self._output_doc_dir_path

            self._log_threaded(f"Loading template: {os.path.basename(template_path)}")
            self._log_threaded(
                f"Writing {len(folder_names)} folder names to Column A..."
            )

            wb = openpyxl.load_workbook(template_path)
            ws = wb.active

            for i, name in enumerate(folder_names, start=2):
                ws.cell(row=i, column=1).value = name

            date_str = datetime.now().strftime("%y%m%d")
            filename = f"{date_str} Gene Seq Log.xlsx"
            save_path = os.path.join(output_dir, filename)

            self._log_threaded(f"Saving file as: {filename}")
            wb.save(save_path)

            self._log_threaded(f"File successfully saved to: {output_dir}")
            self.master.after(
                0, lambda: self.status_label_3.configure(text=f"Generated: {filename}")
            )
            self.master.after(
                0,
                lambda: messagebox.showinfo(
                    "Success", f"Document created successfully:\n{filename}"
                ),
            )
            self._log_threaded("\n--- SG Check Generation Complete ---")

        except Exception as e:
            self._log_threaded(f"\nERROR: {e}")
            self.master.after(0, lambda: messagebox.showerror("Error", str(e)))
        finally:
            self.master.after(
                0,
                lambda: self.btn_generate_sg.configure(
                    state="normal", fg_color="#3C702D"
                ),
            )

    # ─── Utility Methods ────────────────────────────────────────

    @staticmethod
    def _extract_info(filename):
        pattern = re.compile(r"(.+?)_(.+?)_(.+)")
        match = pattern.match(filename)
        if match:
            plasmid_no = match.group(1)
            work_no = match.group(2).strip(".")
            return work_no, plasmid_no
        return None, None


def main():
    ctk.set_appearance_mode("System")
    ctk.set_default_color_theme("blue")
    root = ctk.CTk()
    app = SequencingFileOrganizerApp(root)
    root.mainloop()


if __name__ == "__main__":
    main()
