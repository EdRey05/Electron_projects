"""
Bio Basic Inc. - Canada
Made by: Eduardo Reyes and Fabien Lonjon, Ph.D.
contact: ed5reyes@outlook.com

Version: 1.2
Date: Oct 13, 2026
Notes: Improved file autodetection, naming, and column mapping. Internal
    variables follow the canonical job-log naming (work_no); the
    'input addon' read resolves Work No. / Vector / Insert Sequence /
    Final Vector Sequence by header name first (legacy JobID spellings
    tolerated) with fixed A/E/H/I fallback and logs the fired path.
"""

import os
import sys
import traceback
import threading

import openpyxl

import customtkinter as ctk
from tkinter import filedialog, messagebox


class SequencingReferenceFilesApp:
    """
    Extracts reference sequences from a Primer Log Excel file and
    exports them as individual .txt files.

    Reads Work No. / Vector / Insert Sequence / Final Vector Sequence from
    the 'input addon' sheet, resolved by header name first with the
    fixed A/E/H/I indices as fallback.
    from the "input addon" sheet and creates text files for each sequence.
    """

    def __init__(
        self,
        master,
        primer_log_path=None,
        reference_files_path=None,
    ):
        self.master = master
        master.title("Sequencing Reference Files")
        master.geometry("600x650")

        # --- Validation flags ---
        self.primer_log_valid = False
        self.reference_files_valid = False

        # --- Internal storage ---
        self._primer_log_path = ""
        self._reference_files_path = ""
        self._primer_log_initial_dir = "/"

        # --- Build UI ---
        self._setup_ui()

        # --- Setup initial paths (matches other hub apps pattern) ---
        self.master.after(
            200,
            lambda: self._setup_initial_paths(primer_log_path, reference_files_path),
        )

    def _setup_ui(self):
        """Creates all UI widgets."""

        # --- Primer Log File (Autodetected) ---
        primer_log_frame = ctk.CTkFrame(self.master, fg_color="transparent")
        primer_log_frame.pack(pady=(15, 10), padx=15, fill="x")

        ctk.CTkLabel(
            primer_log_frame,
            text="Primer Log",
            font=ctk.CTkFont(size=18, weight="bold"),
        ).pack(anchor="center")

        entry_frame_pl = ctk.CTkFrame(primer_log_frame, fg_color="transparent")
        entry_frame_pl.pack(fill="x", pady=(5, 0))

        self.entry_primer_log = ctk.CTkEntry(entry_frame_pl, border_width=1)
        self.entry_primer_log.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.button_primer_log = ctk.CTkButton(
            entry_frame_pl,
            text="Select",
            command=self._browse_primer_log,
            width=90,
        )
        self.button_primer_log.pack(side="right")

        self.primer_log_validation_label = ctk.CTkLabel(
            primer_log_frame,
            text="",
            font=ctk.CTkFont(size=14),
            anchor="w",
            justify="left",
        )
        self.primer_log_validation_label.pack(fill="x")

        # --- Reference Files Folder (Autodetected) ---
        reference_files_frame = ctk.CTkFrame(self.master, fg_color="transparent")
        reference_files_frame.pack(pady=(15, 10), padx=15, fill="x")

        ctk.CTkLabel(
            reference_files_frame,
            text="Reference Files Folder",
            font=ctk.CTkFont(size=18, weight="bold"),
        ).pack(anchor="center")

        entry_frame_rf = ctk.CTkFrame(reference_files_frame, fg_color="transparent")
        entry_frame_rf.pack(fill="x", pady=(2, 2))

        self.entry_reference_files = ctk.CTkEntry(entry_frame_rf, border_width=1)
        self.entry_reference_files.pack(side="left", fill="x", expand=True, padx=(0, 10))

        self.button_reference_files = ctk.CTkButton(
            entry_frame_rf,
            text="Select",
            command=self._browse_reference_files,
            width=90,
        )
        self.button_reference_files.pack(side="right")

        self.reference_files_validation_label = ctk.CTkLabel(
            reference_files_frame,
            text="",
            font=ctk.CTkFont(size=14),
            anchor="w",
            justify="left",
        )
        self.reference_files_validation_label.pack(fill="x")

        # --- Export Button ---
        export_button_frame = ctk.CTkFrame(self.master, fg_color="transparent")
        export_button_frame.pack(pady=20, padx=20, fill="x")
        export_button_frame.grid_columnconfigure((0, 2), weight=2)  # Spacers
        export_button_frame.grid_columnconfigure(1, weight=1)  # Button column

        self.button_export = ctk.CTkButton(
            export_button_frame,
            text="Export Sequences",
            command=self._start_processing,
            font=ctk.CTkFont(size=14, weight="bold"),
            height=40,
            fg_color="#3C702D",
            hover_color="#4CAF50",
        )
        self.button_export.configure(state="disabled", fg_color="gray50")
        self.button_export.grid(row=0, column=1, sticky="ew")

        # --- Log Frame ---
        log_frame = ctk.CTkFrame(self.master)
        log_frame.pack(pady=(10, 15), padx=15, fill="both", expand=True)

        ctk.CTkLabel(
            log_frame,
            text="Log",
            font=ctk.CTkFont(size=16, weight="bold"),
        ).pack(anchor="w", padx=10, pady=(5, 2))

        self.log_textbox = ctk.CTkTextbox(
            log_frame,
            activate_scrollbars=True,
            state="disabled",
            fg_color="white",
            text_color="black",
        )
        self.log_textbox.pack(expand=True, fill="both", padx=10, pady=(2, 10))

    # --- Logging methods ---
    def _log(self, message):
        """Appends a message to the log widget (main thread only)."""
        if self.log_textbox and self.log_textbox.winfo_exists():
            self.log_textbox.configure(state="normal")
            self.log_textbox.insert("end", message + "\n")
            self.log_textbox.see("end")
            self.log_textbox.configure(state="disabled")

    def _log_threaded(self, message):
        """Safely log from a worker thread."""
        if self.log_textbox:
            self.master.after(0, self._log, message)

    def _clear_log(self):
        """Clears the log widget."""
        if self.log_textbox and self.log_textbox.winfo_exists():
            self.log_textbox.configure(state="normal")
            self.log_textbox.delete("1.0", "end")
            self.log_textbox.configure(state="disabled")

    # --- Validation methods ---
    def _start_validation_thread(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        expected_sheet,
        validation_flag_name,
        file_display_name,
        next_callback=None,
    ):
        """Starts a background thread to validate a file/folder path."""
        # Reset UI to neutral before starting validation
        self._update_validation_ui(
            entry_widget, validation_label_widget, None, "", validation_flag_name
        )

        thread = threading.Thread(
            target=self._validate_path,
            args=(
                entry_widget,
                validation_label_widget,
                file_path,
                expected_sheet,
                validation_flag_name,
                file_display_name,
                next_callback,
            ),
            daemon=True,
        )
        thread.start()

    def _validate_path(
        self,
        entry_widget,
        validation_label_widget,
        file_path,
        expected_sheet,
        validation_flag_name,
        file_display_name,
        next_callback=None,
    ):
        """
        Checks if a file/folder exists and (for files) contains the expected sheet.
        This method runs in a background thread and schedules UI updates.
        """
        if not file_path:
            return

        # For folders, just check existence (Output folder)
        if validation_flag_name == "reference_files_valid":
            if os.path.isdir(file_path):
                message = "Folder found"
                self.master.after(
                    0,
                    self._update_validation_ui,
                    entry_widget,
                    validation_label_widget,
                    True,
                    message,
                    validation_flag_name,
                    file_display_name,
                    next_callback,
                )
            else:
                message = "Folder not found"
                self.master.after(
                    0,
                    self._update_validation_ui,
                    entry_widget,
                    validation_label_widget,
                    False,
                    message,
                    validation_flag_name,
                    file_display_name,
                    next_callback,
                )
            return

        # For files, check existence and sheet
        if not os.path.exists(file_path):
            message = "File not found."
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                message,
                validation_flag_name,
                file_display_name,
                next_callback,
            )
            return

        try:
            wb = openpyxl.load_workbook(file_path, read_only=True)
            if expected_sheet and expected_sheet not in wb.sheetnames:
                message = f"Sheet '{expected_sheet}' not found."
                self.master.after(
                    0,
                    self._update_validation_ui,
                    entry_widget,
                    validation_label_widget,
                    False,
                    message,
                    validation_flag_name,
                    file_display_name,
                    next_callback,
                )
                return

            # File is valid
            message = "Validated"
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                True,
                message,
                validation_flag_name,
                file_display_name,
                next_callback,
            )
        except Exception as e:
            message = "Not a valid Excel file."
            self.master.after(
                0,
                self._update_validation_ui,
                entry_widget,
                validation_label_widget,
                False,
                message,
                validation_flag_name,
                file_display_name,
                next_callback,
            )

    def _update_validation_ui(
        self,
        entry_widget,
        validation_label_widget,
        is_valid,
        message,
        validation_flag_name,
        file_display_name=None,
        next_callback=None,
    ):
        """
        Updates the UI to reflect the validation status of a file/folder.
        This method is always run on the main UI thread.
        """
        # Update the validation status flag
        setattr(self, validation_flag_name, is_valid)

        # Store the path if valid
        if is_valid is True:
            if validation_flag_name == "primer_log_valid":
                self._primer_log_path = entry_widget.get()
            elif validation_flag_name == "reference_files_valid":
                self._reference_files_path = entry_widget.get()

            entry_widget.configure(border_color="dark green", border_width=2)
            validation_label_widget.configure(text=f"✔ {message}", text_color="green")
        elif is_valid is False:
            entry_widget.configure(border_color="#D32F2F", border_width=2)
            validation_label_widget.configure(
                text=f"✖ Invalid: {message}", text_color="#D32F2F"
            )
        else:  # is_valid is None (resetting to neutral)
            entry_widget.configure(
                border_width=1,
                border_color=ctk.ThemeManager.theme["CTkEntry"]["border_color"],
            )
            validation_label_widget.configure(
                text="", text_color=ctk.ThemeManager.theme["CTkLabel"]["text_color"]
            )

        self._check_all_validations()

        if next_callback:
            next_callback()

    def _check_all_validations(self):
        """Checks if all inputs are valid and enables/disables the Export button."""
        all_valid = self.primer_log_valid and self.reference_files_valid
        if all_valid:
            self.button_export.configure(state="normal", fg_color="#3C702D")
        else:
            self.button_export.configure(state="disabled", fg_color="gray50")

    def _setup_initial_paths(self, primer_log_path, reference_files_path):
        """Pre-fills entry widgets with provided paths and then starts sequential validation."""
        # --- Phase 1: Autofill all paths immediately ---
        if primer_log_path:
            # Check if it's a file or directory
            if os.path.isfile(primer_log_path):
                self.entry_primer_log.delete(0, "end")
                self.entry_primer_log.insert(0, primer_log_path)
            elif os.path.isdir(primer_log_path):
                self._primer_log_initial_dir = primer_log_path

        if reference_files_path:
            self.entry_reference_files.delete(0, "end")
            self.entry_reference_files.insert(0, reference_files_path)

        # --- Phase 2: Start the validation chain ---
        self._start_sequential_validation()

    def _start_sequential_validation(self):
        """Starts the validation for any pre-filled paths in sequence."""
        # Get paths from the entry widgets now that they are filled
        primer_log_path = self.entry_primer_log.get()
        reference_files_path = self.entry_reference_files.get()

        # Define chained callbacks

        # 2. Reference Files Folder (Last)
        def validate_reference_files():
            if reference_files_path:
                self._start_validation_thread(
                    self.entry_reference_files,
                    self.reference_files_validation_label,
                    reference_files_path,
                    None,  # No sheet check for folders
                    "reference_files_valid",
                    "Reference Files Folder",
                )

        # 1. Primer Log (First)
        if primer_log_path:
            self._start_validation_thread(
                self.entry_primer_log,
                self.primer_log_validation_label,
                primer_log_path,
                "input addon",
                "primer_log_valid",
                "Primer Log",
                next_callback=validate_reference_files,
            )
        else:
            validate_reference_files()  # Skip to next if path is missing

    # --- Browse methods ---
    def _browse_primer_log(self):
        filename = filedialog.askopenfilename(
            title="Select Primer Log File",
            initialdir=self._primer_log_initial_dir,
            filetypes=[("Excel files", "*.xlsx *.xls"), ("All files", "*.*")],
        )
        if filename:
            self.entry_primer_log.delete(0, "end")
            self.entry_primer_log.insert(0, filename)
            self._start_validation_thread(
                self.entry_primer_log,
                self.primer_log_validation_label,
                filename,
                "input addon",
                "primer_log_valid",
                "Primer Log",
            )

    def _browse_reference_files(self):
        folder = filedialog.askdirectory(title="Select Reference Files Folder")
        if folder:
            self.entry_reference_files.delete(0, "end")
            self.entry_reference_files.insert(0, folder)
            self._start_validation_thread(
                self.entry_reference_files,
                self.reference_files_validation_label,
                folder,
                None,
                "reference_files_valid",
                "Reference Files Folder",
            )

    # --- Processing methods ---
    def _start_processing(self):
        """Starts the processing logic in a background thread."""
        self.button_export.configure(state="disabled", fg_color="gray50")
        self._clear_log()
        self._log("Processing started...")

        thread = threading.Thread(target=self._processing_thread, daemon=True)
        thread.start()

    def _processing_thread(self):
        """The main processing logic that runs in a background thread."""
        try:
            self._log_threaded(f"Loading Primer Log: {os.path.basename(self._primer_log_path)}")

            # Load the input addon sheet
            wb = openpyxl.load_workbook(self._primer_log_path, read_only=True, data_only=True)
            
            sheet_name = "input addon"
            if sheet_name not in wb.sheetnames:
                self._log_threaded(f"ERROR: '{sheet_name}' sheet not found in the file.")
                self.master.after(
                    0,
                    lambda: messagebox.showerror("Error", f"Sheet '{sheet_name}' not found in the file.")
                )
                self.master.after(
                    0, lambda: self.button_export.configure(state="normal", fg_color="#3C702D")
                )
                return

            ws = wb[sheet_name]

            # Count sequences to export
            sequences_exported = 0
            rows_processed = 0

            self._log_threaded(f"Reference folder: {self._reference_files_path}")
            self._log_threaded("")

            # Read header row (Row 1) to determine indices
            rows_iter = ws.iter_rows(values_only=True)
            headers = next(rows_iter, None)
            
            if not headers:
                self._log_threaded("ERROR: File is empty.")
                return

            # Map normalized column names to indices (both sides stripped +
            # lowercased so surrounding whitespace never breaks a match).
            import re

            def norm_header(h):
                return re.sub(r"\s+", " ", str(h)).strip().lower() if h is not None else ""

            header_map = {norm_header(h): i for i, h in enumerate(headers) if h}

            def resolve_column(candidates, fallback_idx):
                """Canonical-first candidates; legacy spellings tolerated.
                Falls back to the historic fixed index when no candidate
                header matches. Returns (idx, matched_by_name)."""
                for cand in candidates:
                    key = re.sub(r"\s+", " ", str(cand)).strip().lower()
                    if key in header_map:
                        return header_map[key], True
                return fallback_idx, False

            # Fixed fallback indices: A=0, E=4, H=7, I=8.
            idx_work_no, name_work_no = resolve_column(
                ("Work No.", "JobID", "Job ID"), 0)
            idx_vector, name_vector = resolve_column(("Vector",), 4)
            idx_insert_sequence, name_insert = resolve_column(
                ("Insert Sequence", "Insert Seq"), 7)
            idx_final_vector_sequence, name_final = resolve_column(
                ("Final Vector Sequence", "Final Vector"), 8)

            def col_letter(i):
                from openpyxl.utils import get_column_letter
                return get_column_letter(i + 1)

            col_report = ", ".join(
                f"{label}={col_letter(idx)}"
                f"{' (name)' if used_name else ' (fixed fallback)'}"
                for label, idx, used_name in (
                    ("Work No.", idx_work_no, name_work_no),
                    ("Vector", idx_vector, name_vector),
                    ("Insert Sequence", idx_insert_sequence, name_insert),
                    ("Final Vector Sequence", idx_final_vector_sequence, name_final),
                )
            )
            if name_work_no and name_vector and name_insert and name_final:
                self._log_threaded(f"'input addon': columns resolved by name ({col_report})")
            elif not (name_work_no or name_vector or name_insert or name_final):
                self._log_threaded(
                    f"'input addon': header row not recognized; falling back to "
                    f"fixed columns (Work No.=A, Vector=E, Insert Sequence=H, "
                    f"Final Vector Sequence=I)"
                )
            else:
                self._log_threaded(
                    f"'input addon': partial header match; fixed fallback used "
                    f"for missing columns ({col_report})"
                )

            # Process data rows
            for row in rows_iter:
                # Helper to safely get value by index
                def get_val(idx):
                    return str(row[idx]).strip() if idx < len(row) and row[idx] is not None else ""

                work_no = get_val(idx_work_no)
                vector = get_val(idx_vector)
                insert_sequence = get_val(idx_insert_sequence)
                final_vector_sequence = get_val(idx_final_vector_sequence)

                # Skip empty rows
                if not work_no:
                    continue

                rows_processed += 1

                # File 1: {Work No.}.txt with Insert Sequence
                if insert_sequence:
                    file_name = f"{work_no}.txt"
                    file_path = os.path.join(self._reference_files_path, file_name)
                    with open(file_path, 'w', encoding='utf-8') as f:
                        f.write(insert_sequence)
                    self._log_threaded(f"  ✓ Created: {file_name}")
                    sequences_exported += 1

                # File 2: {Work No.}+{Vector}.txt with Final Vector Sequence (if present)
                if final_vector_sequence:
                    file_name = f"{work_no}+{vector}.txt"
                    file_path = os.path.join(self._reference_files_path, file_name)
                    with open(file_path, 'w', encoding='utf-8') as f:
                        f.write(final_vector_sequence)
                    self._log_threaded(f"  ✓ Created: {file_name}")
                    sequences_exported += 1

            wb.close()

            self._log_threaded("")
            self._log_threaded(f"--- Processing Finished ---")
            self._log_threaded(f"Rows processed: {rows_processed}")
            self._log_threaded(f"Files created: {sequences_exported}")

            self.master.after(
                0,
                lambda: messagebox.showinfo(
                    "Success",
                    f"Successfully exported {sequences_exported} sequence files.",
                ),
            )

        except Exception as e:
            self._log_threaded(f"\nERROR: {traceback.format_exc()}")
            self.master.after(0, messagebox.showerror, "Error", str(e))
        finally:
            self.master.after(
                0, lambda: self.button_export.configure(state="normal", fg_color="#3C702D")
            )


def main():
    """Standalone entry point when not launched from the hub."""
    ctk.set_appearance_mode("System")
    ctk.set_default_color_theme("blue")

    # Parse arguments for standalone execution logic matching Hub's format
    kwargs = {}
    for arg in sys.argv[1:]:
        if arg.startswith("--primer_log_path="):
            kwargs["primer_log_path"] = arg.split("=", 1)[1]
        elif arg.startswith("--reference_files_path="):
            kwargs["reference_files_path"] = arg.split("=", 1)[1]

    root = ctk.CTk()
    app = SequencingReferenceFilesApp(root, **kwargs)
    root.mainloop()


if __name__ == "__main__":
    main()