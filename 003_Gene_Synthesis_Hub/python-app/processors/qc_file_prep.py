"""Headless adapters for the seven reference QC File Prep workflows.

Upstream processing stays in Code/. Only UI state, dialogs, paths, staging,
Word lifecycle and publication are adapted here.
"""
import configparser
from pathlib import Path
from types import SimpleNamespace
import math
import os
import re
import shutil
import tempfile
import uuid
import zipfile
from processors.common import load_upstream,validate_archive

WORKFLOWS={
 'inhouse-new':('inhouse_newnames','InHouseNewNames'),
 'inhouse-old':('inhouse_oldnames','InHouseOldNames'),
 'internal':('internal_renaming','InternalRenaming'),
 'wgk':('wgk_correction','WgkCorrection'),
 'resuspension':('resuspension_prep','ResuspensionPrep'),
 'batch-post':('batch_post_processing','BatchPostProcessing'),
 'scratch':('inhouse_qc_new','InHouseQcNew'),
}

class Value:
    def __init__(self,value='',**_):self.value=value
    def get(self,*_):return self.value
    def set(self,value):self.value=value

def safe_component(text):
    if not text or re.search(r'[<>:"/\\|?*\x00-\x1f]',str(text)) or str(text).endswith(('.', ' ')) or str(text) in ('.','..'):
        raise ValueError('Invalid filename component: '+str(text))
    if re.match(r'^(CON|PRN|AUX|NUL|COM[1-9]|LPT[1-9])(\.|$)',str(text),re.I):raise ValueError('Reserved filename: '+str(text))
    return str(text)

def native_header(doc_path,log):
    """Keep original drawing geometry; change fonts without requiring Word."""
    from docx import Document
    doc=Document(doc_path)
    for paragraph in doc.paragraphs:
        for run in paragraph.runs:run.font.name='Calibri'
    for table in doc.tables:
        for row in table.rows:
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:run.font.name='Calibri'
    doc.save(doc_path)
    log('[REVIEW] Header shapes preserved. Check report header alignment in the finished DOCX.')

def run(payload,root,log,ask):
    name=payload.get('workflow')
    if name not in WORKFLOWS:raise ValueError('Unknown QC File Prep workflow.')
    output=Path(payload['outputFolder']).resolve()
    if not output.is_dir():raise ValueError('Choose an existing output folder.')
    warnings=[];words=[]
    def record(message):
        text=str(message);log(text)
        if re.search(r'\[(ERROR|FATAL|FAIL|WARN|WARNING|REVIEW|SKIP)\]',text):warnings.append(text)
    def notify(title,message,*args,**kwargs):record('[WARNING] '+str(title)+': '+str(message))
    def error(title,message,*args,**kwargs):raise ValueError(str(title)+': '+str(message))
    module_name,class_name=WORKFLOWS[name]
    module=load_upstream('gsh_qc_'+module_name,root/'Code/QC_file_prep/workflows'/(module_name+'.py'))
    module.ctk=SimpleNamespace(StringVar=Value,BooleanVar=Value)
    module.messagebox=SimpleNamespace(showerror=error,showwarning=notify,showinfo=notify,askyesno=lambda title,message,**_:ask(title,message,['cancel','ok'])['choice']=='ok')
    def private_word(*_):
        if not payload.get('allowWord'):raise ValueError('Legacy DOC requires Word. Use DOCX or enable the optional installed-Word compatibility setting.')
        import win32com.client
        word=win32com.client.DispatchEx('Word.Application');words.append(word);word.Visible=False;word.DisplayAlerts=0
        return word
    if hasattr(module,'win32'):module.win32=SimpleNamespace(Dispatch=private_word,DispatchEx=private_word)
    try:
        with tempfile.TemporaryDirectory(prefix='gsh-qc-prep-') as temporary:
            workspace=Path(temporary);inputs=workspace/'input';inputs.mkdir();private_output=workspace/'scratch-output';private_output.mkdir()
            paths={**payload.get('paths',{}),'QC_Input_Folder':str(inputs),'QC_Output_Folder':str(private_output),'Default_Output_QC_Folder':str(private_output)}
            paths['Default_Gel_Images_Folder']=paths.get('QC_Gel_Images','')
            config=configparser.ConfigParser();config['DEFAULT']={'Max_Img_Width_Inches':str(payload.get('maxWidth',2.3)),'Max_Img_Height_Inches':str(payload.get('maxHeight',2.8))}
            for key in ('Max_Img_Width_Inches','Max_Img_Height_Inches'):
                value=float(config['DEFAULT'][key])
                if not math.isfinite(value) or value<=0 or value>20:raise ValueError('Image dimensions must be positive and at most 20 inches.')
            app=SimpleNamespace(log=record,config=config,input_folder=Value(str(inputs)),hub_paths={**paths,'Gel_Ladder':paths.get('QC_Gel_Ladder','')},make_path_from_config=lambda section,key:paths.get(key,''),process_button=SimpleNamespace(configure=lambda **_:None),master=SimpleNamespace(after=lambda delay,fn,*args:fn(*args)))
            processor=getattr(module,class_name)(app)
            # These copy/cleanup options belonged to the Tk host. Private staging
            # replaces them; originals and network destinations are never changed.
            for attribute in ('copy_output_var','delete_intermediate_var'):
                if hasattr(processor,attribute):getattr(processor,attribute).set(False)
            if name!='scratch':
                source=Path(payload['inputFolder']).resolve()
                if not source.is_dir():raise ValueError('Choose the folder containing the source ZIPs.')
                archives=sorted(p for p in source.iterdir() if p.is_file() and p.suffix.lower()=='.zip')
                if not archives:raise ValueError('No ZIP files found.')
                for archive in archives:
                    safe_component(archive.stem);validate_archive(archive)
                    # Reject unsupported binary reports before any processing.
                    with zipfile.ZipFile(archive) as zipped:
                        if any(Path(n).suffix.lower()=='.doc' for n in zipped.namelist()) and not payload.get('allowWord'):
                            raise ValueError(archive.name+' contains legacy DOC. DOCX works without Office; binary DOC needs installed Word compatibility.')
                    shutil.copy2(archive,inputs/archive.name)
            if name in ('inhouse-new','inhouse-old'):
                if name=='inhouse-new':
                    processor.rename_from_excel_var.set(bool(payload.get('renameFromLog',True)))
                    processor.fix_header_formatting=lambda path:native_header(path,record)
                if not processor.validate_inputs():raise ValueError('Invalid workflow inputs.')
            elif name=='internal':
                processor.replacement_file.set(payload.get('replacementFile',''));processor.load_replacement_strings()
                if not processor.replacement_map:raise ValueError('Load a non-empty tab-separated replacement map.')
                # A map must never introduce a path or collapse distinct names.
                for old,new in processor.replacement_map.items():
                    if any(c in old+new for c in '/\\\x00'):raise ValueError('Replacement values cannot contain path separators.')
                for archive in inputs.glob('*.zip'):
                    with zipfile.ZipFile(archive) as zipped:
                        renamed=[]
                        for entry in zipped.namelist():
                            for old,new in processor.replacement_map.items():entry=entry.replace(old,new)
                            normalized=entry.rstrip('/').casefold()
                            if normalized in renamed:raise ValueError('Replacement map produces duplicate archive paths.')
                            renamed.append(normalized)
                processor.remove_qc_table_var.set(bool(payload.get('removeQcTable')))
            elif name=='wgk':
                processor.input_mode.set(payload.get('quantityMode','single'))
                processor.plasmid_qty.set(payload.get('quantity','10 µg'));processor.batch_file_path.set(payload.get('replacementFile',''))
                if not processor.validate_inputs():raise ValueError('Invalid quantity inputs.')
            elif name=='batch-post':
                def assignments(ready,input_folder=None):
                    result={}
                    for entry in ready:
                        fields=[{'key':'vector','label':entry['vector']+' vector FASTA','type':'file','extension':'.fasta','value':payload.get('vectorFasta','')}]
                        fields += [{'key':f'od{i}','label':label,'value':str((payload.get('odValues') or ['','','',''])[i])} for i,label in enumerate(['OD260','OD280','OD260/OD280','µg/µL'])]
                        while True:
                            reply=ask('BATCH vector and OD',entry['zip_name']+'\nConfirm vector FASTA and all four OD values.',['cancel','skip','proceed'],fields)
                            if reply['choice']=='skip':result[entry['zip_stem']]={'skipped':True};break
                            values=reply.get('data') or {};vector=Path(values.get('vector',''))
                            try:
                                if not vector.is_file() or vector.suffix.lower() not in ('.fasta','.fa','.fna'):raise ValueError('Choose a vector FASTA file.')
                                if vector.stem.casefold()!=entry['vector'].casefold():raise ValueError('FASTA filename must match vector '+entry['vector'])
                                content=vector.read_text(encoding='utf-8-sig').strip().splitlines()
                                if not content or content[0]!='>'+vector.stem or len([v for v in content if v.startswith('>')])!=1:raise ValueError('FASTA must have one header and one sequence.')
                                sequence=''.join(content[1:]).replace(' ','').upper()
                                if not sequence or re.search('[^ACGTRYSWKMBDHVN]',sequence):raise ValueError('FASTA has annotations or invalid DNA.')
                                od=[str(values.get(f'od{i}','')).strip() for i in range(4)]
                                if not all(math.isfinite(float(v)) and float(v)>=0 for v in od):raise ValueError('Enter four non-negative OD measurements.')
                                for key in ('work_no','bbi_id','gene_name','vector'):safe_component(entry[key])
                                result[entry['zip_stem']]={'vector_fasta_path':str(vector),'od_values':od};break
                            except (ValueError,OSError) as exc:
                                record('[WARNING] '+str(exc));fields=[{**f,'value':values.get(f['key'],'')} for f in fields]
                    return result
                processor._show_vector_selection_dialog=assignments
                processor._blocking_dialog=lambda message:record('[ERROR] '+message)
            elif name=='scratch':
                module.TEMPLATE_DIR=str(root/'Templates/QC_zip_template')
                ids=list(dict.fromkeys(v.strip().upper() for v in payload.get('workNos',[]) if v.strip()))
                for value in ids:safe_component(value)
                if not ids:raise ValueError('Enter at least one Work No.')
                processor.txt_ids=Value('\n'.join(ids))
                def alignments(ready):
                    result={};alignment=Path(payload.get('alignmentFolder',''))
                    for entry in ready:
                        value=entry['work_no'];row=entry['row']
                        for key in ('BBI ID','Vector','Plasmid No.','Tube ID'):safe_component(str(row[key]).strip())
                        def locate(extension):
                            matches=[p for p in alignment.glob('*') if p.is_file() and p.suffix.lower()==extension and re.match(re.escape(value)+r'(?=[^A-Za-z0-9]|$)',p.stem,re.I)] if alignment.is_dir() else []
                            return str(matches[0]) if len(matches)==1 else ''
                        fields=[{'key':'pdf','label':value+' alignment PDF','type':'file','extension':'.pdf','value':locate('.pdf')},{'key':'sqd','label':value+' alignment SQD','type':'file','extension':'.sqd','value':locate('.sqd')}]
                        while True:
                            reply=ask('Alignment files',value+' — '+entry['vector'],['cancel','skip','proceed'],fields)
                            if reply['choice']=='skip':result[value]={'skipped':True};break
                            values=reply.get('data') or {}
                            if all(Path(values.get(k,'')).is_file() and Path(values[k]).suffix.lower()=='.'+k for k in ('pdf','sqd')):result[value]=values;break
                            record('[WARNING] Choose an existing PDF and SQD.');fields=[{**f,'value':values.get(f['key'],'')} for f in fields]
                    return result
                processor._show_alignment_dialog=alignments
            processor.run_processing_task()
            products=[]
            for directory in workspace.iterdir():
                if directory.is_dir() and ('output_zip' in directory.name or directory==private_output):products.extend(directory.glob('*.zip'))
            if name=='resuspension':products.extend(inputs.glob('Resuspension_Prep_Data_*.docx'))
            if not products:raise ValueError('No output was produced. Check the processing log for skipped or failed records.')
            for product in products:
                product.resolve().relative_to(workspace)
                if product.suffix.lower()=='.zip':validate_archive(product)
            # One unique run directory prevents overwriting earlier exports.
            destination=output/('GSH_'+name+'_'+uuid.uuid4().hex[:12]);destination.mkdir()
            files=[]
            for product in products:
                target=destination/product.name
                if target.exists():raise ValueError('Duplicate output name: '+product.name)
                with product.open('rb') as source,target.open('xb') as output_file:shutil.copyfileobj(source,output_file)
                files.append({'name':target.name,'path':str(target)})
            return {'files':files,'warnings':warnings,'outputFolder':str(destination)}
    finally:
        for word in words:
            try:word.Quit()
            except Exception:pass
